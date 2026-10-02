"""Generated PO regressions for missing catalog information and warning notes."""

from __future__ import annotations

import datetime
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

import main
from test_po_variants import TEMPLATE, combine, make_barcode_catalog, source_row


def complete_entry(**changes) -> dict:
    entry = {
        "code": "MC-510",
        "description": "Catalog red faucet",
        "brand": "Donmark",
        "material": "Steel",
        "weight": 1.5,
        "carton": 12,
        "barcode": "000123",
        "picture_color": "red",
    }
    entry.update(changes)
    return entry


def is_yellow(cell) -> bool:
    color = cell.fill.fgColor
    return cell.fill.patternType == "solid" and color.type == "rgb" and color.rgb[-6:] == "FFFF00"


class MissingCatalogTests(unittest.TestCase):
    def generate(
        self,
        output: Path,
        source: list[dict],
        entries: list[dict],
        *,
        catalog_filename: str | None = None,
    ):
        catalog_path = output / "catalog.xlsx"
        make_barcode_catalog(catalog_path, entries)
        rows = combine([], source)
        with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
            path = main.generate_po_from_combined(
                rows, "A0029", datetime.date(2026, 10, 2), 6,
                str(TEMPLATE), str(catalog_path), str(output / "missing_vendors.xlsx"),
                4, 7, catalog_filename=catalog_filename,
            )
        return openpyxl.load_workbook(path)["PO"]

    def assert_missing_quantity_formulas_are_guarded(self, po, row: int) -> None:
        columns = main.get_po_col_map(po)
        qpc_reference = f'{get_column_letter(columns["QTY PER CARTON"])}{row}'
        self.assertIsNone(po.cell(row, columns["QTY PER CARTON"]).value)
        for header in ("CARTONS", "GREEN", "TOTAL QTY (ORDER)", "AMOUNT (CNY)"):
            with self.subTest(header=header):
                cell = po.cell(row, columns[header])
                self.assertTrue(cell.value.startswith("=IF("), cell.value)
                self.assertIn('""', cell.value)
                self.assertIn(qpc_reference, cell.value)
                self.assertTrue(is_yellow(cell))

    def test_absent_item_exports_source_identity_and_blank_yellow_catalog_fields(self) -> None:
        description = "ก๊อกอ่างล้างหน้า สีแดง ขนาด 1/2 /OEM"
        source = [source_row(description, code="MISSING-ITEM", barcode="0000999", sales=6, stock=1)]
        filename = "ใบรายการสินค้า 02 ต.ค. 2569.xlsx"
        with tempfile.TemporaryDirectory() as tmp:
            po = self.generate(
                Path(tmp), source, [complete_entry(code="OTHER-ITEM")],
                catalog_filename=filename,
            )

            columns = main.get_po_col_map(po)
            self.assertEqual(po["A9"].value, "MISSING-ITEM")
            self.assertEqual(po["C9"].value, description)
            self.assertEqual(po["Y9"].value, "0000999")
            self.assertEqual(po["Y9"].data_type, "s")
            self.assertEqual(po["Y9"].number_format, "@")
            for header in ("BRAND", "MATERIAL", "Weight", "QTY PER CARTON"):
                self.assertIsNone(po.cell(9, columns[header]).value)
            self.assertFalse(any(image.anchor._from.col == 1 for image in po._images))
            self.assertTrue(all(is_yellow(po.cell(9, col)) for col in range(1, 26)))
            self.assertEqual(po["D9"].value, f'ไม่มีรายละเอียดสินค้าตัวนี้ อัปเดต "{filename}"')
            self.assertNotIn('"catalog.xlsx"', po["D9"].value)
            self.assert_missing_quantity_formulas_are_guarded(po, 9)

    def test_absent_item_returns_no_catalog_match_even_when_other_barcodes_exist(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "catalog.xlsx"
            make_barcode_catalog(path, [complete_entry(code="OTHER-ITEM")])
            catalog = main.build_catalog_map(str(path), "A0029")

            selected = main.resolve_catalog_variant(
                catalog, "MISSING-ITEM", "GREEN source description", 1,
                barcode="0000999",
            )

            self.assertEqual(selected, {})

    def test_missing_and_invalid_carton_quantities_export_as_blank_with_exact_note(self) -> None:
        for carton in (None, "", "not a number", 0, -12, "NaN", "inf"):
            with self.subTest(carton=carton), tempfile.TemporaryDirectory() as tmp:
                po = self.generate(
                    Path(tmp), [source_row("Red", barcode="000123", sales=6)],
                    [complete_entry(carton=carton)],
                )

                self.assertEqual(po["D9"].value, 'ไม่มี "carton" ใน "catalog.xlsx"')
                self.assertIsNone(po["H9"].value)
                self.assertTrue(is_yellow(po["H9"]))
                self.assertTrue(is_yellow(po["D9"]))
                self.assertFalse(is_yellow(po["E9"]))
                self.assertEqual(po["E9"].value, "Donmark")
                self.assert_missing_quantity_formulas_are_guarded(po, 9)

    def test_complete_row_after_missing_first_row_does_not_inherit_warning_fill(self) -> None:
        sources = [
            source_row("Missing product", code="A-MISSING", barcode="0000999", sales=6),
            source_row("Complete product", code="Z-COMPLETE", barcode="000123", sales=6),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            po = self.generate(Path(tmp), sources, [complete_entry(code="Z-COMPLETE")])

            self.assertEqual(po["A9"].value, "A-MISSING")
            self.assertEqual(po["A10"].value, "Z-COMPLETE")
            self.assertTrue(all(is_yellow(po.cell(9, col)) for col in range(1, 26)))
            self.assertFalse(any(is_yellow(po.cell(10, col)) for col in range(1, 26)))
            self.assertIn(po["D10"].value, (None, ""))
            self.assertEqual(po["H10"].value, 12)
            self.assertEqual(po["I10"].value, "=ROUND((R10-V10)/H10,0)")
            self.assertEqual(po["J10"].value, "=I10*H10")
            self.assertEqual(po["L10"].value, "=J10+K10")
            self.assertEqual(po["O10"].value, "=M10*L10")

            total_row, _ = main.find_label_cell(po, "TOTAL AMOUNT CNY")
            for column in ("I", "L", "O"):
                formula = po[f"{column}{total_row}"].value
                self.assertTrue(formula.startswith("=IF("), formula)
                self.assertIn(f"COUNT({column}9:{column}10)=2", formula)
                self.assertIn(f"SUM({column}9:{column}10)", formula)

    def test_individual_missing_catalog_fields_have_specific_notes_and_yellow_cells(self) -> None:
        cases = (
            ("description", None, "รายละเอียดสินค้า", "C"),
            ("picture_color", None, "รูปสินค้า", "B"),
            ("brand", None, "ยี่ห้อ", "E"),
            ("material", None, "วัสดุ", "F"),
            ("weight", None, "น้ำหนัก", "G"),
        )
        filename = "รายละเอียดสินค้า รุ่นใหม่.xlsx"
        for field, value, label, column in cases:
            with self.subTest(field=field), tempfile.TemporaryDirectory() as tmp:
                po = self.generate(
                    Path(tmp), [source_row("GREEN full product name", barcode="000123", sales=6)],
                    [complete_entry(**{field: value})], catalog_filename=filename,
                )

                self.assertEqual(po["D9"].value, f'ไม่มี "{label}" ใน "{filename}"')
                self.assertTrue(is_yellow(po[f"{column}9"]))
                self.assertTrue(is_yellow(po["D9"]))
                self.assertEqual(po["C9"].value, "GREEN full product name")
                if column in ("E", "F", "G"):
                    self.assertIsNone(po[f"{column}9"].value)
                self.assertEqual(po["H9"].value, 12)
                self.assertFalse(is_yellow(po["H9"]))
                self.assertEqual(po["I9"].value, "=ROUND((R9-V9)/H9,0)")

    def test_missing_barcode_alone_does_not_create_missing_information_warning(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            po = self.generate(
                Path(tmp), [source_row("Red faucet", barcode="", sales=6)],
                [complete_entry(description="Red faucet", barcode="")],
            )

            self.assertIn(po["D9"].value, (None, ""))
            self.assertIsNone(po["Y9"].value)
            self.assertFalse(any(is_yellow(po.cell(9, col)) for col in range(1, 26)))

    def test_note_insertion_preserves_header_positions_and_shifted_normal_formulas(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            po = self.generate(
                Path(tmp), [source_row("Red", barcode="000123", sales=6, stock=1, on_order=2)],
                [complete_entry()],
            )

            columns = main.get_po_col_map(po)
            self.assertEqual(po["H6"].value, "A0029")
            self.assertEqual(po["C8"].value, "GOODS DESCRIPTION")
            self.assertEqual(po["D8"].value, "หมายเหตุ")
            self.assertEqual(columns["BRAND"], 5)
            self.assertEqual(columns["QTY PER CARTON"], 8)
            self.assertEqual(columns["BARCODE"], 25)
            self.assertEqual(columns["BARCODE"], max(columns.values()))
            self.assertEqual(po["Y8"].value, "BARCODE")
            self.assertIn("$Y", str(po.print_area))
            expected = {
                "Q9": "=P9*4", "R9": "=P9*7", "V9": "=S9+T9+U9",
                "I9": "=ROUND((R9-V9)/H9,0)", "J9": "=I9*H9",
                "L9": "=J9+K9", "O9": "=M9*L9",
            }
            for address, formula in expected.items():
                self.assertEqual(po[address].value, formula)
            total_row, _ = main.find_label_cell(po, "TOTAL AMOUNT CNY")
            for column in ("I", "L", "O"):
                self.assertEqual(po[f"{column}{total_row}"].value, f"=SUM({column}9:{column}9)")

    def test_streamlit_entry_propagates_original_uploaded_catalog_filename(self) -> None:
        filename = "บัญชีรูปสินค้า Supplier 2 ตุลาคม.xlsx"
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog.xlsx"
            make_barcode_catalog(catalog_path, [complete_entry(carton=None)])
            express_path = output / "Express_G.xlsx"
            express_path.touch()
            export_all = main.export_vendor_all_items_excel
            source = pd.DataFrame([source_row("Red", barcode="000123", sales=6)])
            with (
                patch.object(main, "PO_OUTPUT_FOLDER", str(output)),
                patch.object(main, "parse_express_file", return_value=(source, {"months": 3})),
                patch.object(main, "export_vendor_all_items_excel", side_effect=lambda rows, vendor_code: export_all(rows, vendor_code, out_folder=str(output))),
            ):
                result = main.generate_po_streamlit(
                    express_asia_path="", express_green_path=str(express_path),
                    catalog_path=str(catalog_path), vendor_info_path=str(output / "missing_vendors.xlsx"),
                    template_path=str(TEMPLATE), vendor_code="A0029",
                    po_date=datetime.date(2026, 10, 2), rate_thb_per_cny=6,
                    min_factor=4, max_factor=7, catalog_filename=filename,
                )

            self.assertEqual(result["count_all"], 1)
            self.assertEqual(result["count_filtered"], 1)
            po = openpyxl.load_workbook(result["po_filtered"])["PO"]
            self.assertEqual(po["D9"].value, f'ไม่มี "carton" ใน "{filename}"')


if __name__ == "__main__":
    unittest.main()
