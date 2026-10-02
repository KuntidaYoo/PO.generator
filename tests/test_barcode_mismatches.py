"""Cross-report barcode mismatches retain GREEN identity and visible warnings."""

from __future__ import annotations

import datetime
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

import main
from test_po_variants import TEMPLATE, combine, make_barcode_catalog, source_row


WARNING = "barcode ทั้ง2ไฟล์ไม่ตรงกัน"


def catalog_entry(**changes) -> dict:
    entry = {
        "code": "MC-510", "description": "Catalog red faucet", "barcode": "000222",
        "brand": "GREEN catalog brand", "material": "Steel", "weight": 1.5,
        "carton": 12, "picture_color": "red",
    }
    entry.update(changes)
    return entry


class BarcodeMismatchTests(unittest.TestCase):
    def matching_mismatch(self):
        return combine(
            [source_row("ก็อกแฟนซี สีแดง-MR", barcode="000111", sales=6, stock=2, on_order=1, yuan=10)],
            [source_row("ก๊อกแฟนซี สีแดง /OEM", barcode="000222", sales=9, stock=3, on_order=4, yuan=11)],
        )

    def generate_po(self, rows, output: Path, entries: list[dict], filename="catalog.xlsx"):
        path = output / "catalog.xlsx"
        make_barcode_catalog(path, entries)
        with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
            generated = main.generate_po_from_combined(
                rows, "A0029", datetime.date(2026, 10, 2), 6,
                str(TEMPLATE), str(path), str(output / "missing_vendors.xlsx"),
                4, 7, catalog_filename=filename,
            )
        return openpyxl.load_workbook(generated)["PO"]

    def test_same_product_different_barcodes_combines_totals_and_preserves_both_identifiers(self) -> None:
        rows = self.matching_mismatch()

        self.assertEqual(len(rows), 1)
        row = rows.iloc[0]
        self.assertEqual(row["barcode"], "000222")
        self.assertEqual(row["catalog_match_barcode"], "000222")
        self.assertEqual(row["barcode_ASIA"], "000111")
        self.assertEqual(row["barcode_GREEN"], "000222")
        self.assertTrue(row["barcode_mismatch"])
        self.assertEqual(row["หมายเหตุ"], WARNING)
        self.assertEqual(row["รายละเอียดสินค้า"], "ก๊อกแฟนซี สีแดง /OEM")
        self.assertEqual(row["ยอดขาย_TOTAL"], 15)
        self.assertEqual(row["ON_ORDER_TOTAL"], 5)
        self.assertEqual(row["TOTAL_QTY_NUM"], 10)
        self.assertEqual(row["USE_MONTH"], 5)
        self.assertEqual((row["MIN_NUM"], row["MAX_NUM"]), (20, 35))
        self.assertEqual((row["หยวน_ASIA"], row["หยวน_GREEN"], row["หยวน"]), (10, 11, 11))

    def test_item_code_punctuation_and_compatible_truncated_wording_still_merge(self) -> None:
        rows = combine(
            [source_row("ฝารองนั่งวงรี", code="K-1900", barcode="000111", sales=2)],
            [source_row("ฝารองนั่งรุ่นวงรี สีขาว", code="K1900", barcode="000222", sales=3)],
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows.iloc[0]["รหัสสินค้า"], "K1900")
        self.assertEqual(rows.iloc[0]["barcode"], "000222")
        self.assertEqual(rows.iloc[0]["หมายเหตุ"], WARNING)

    def test_same_source_repeated_barcode_quantities_still_aggregate_before_matching(self) -> None:
        rows = combine(
            [source_row("Red", barcode="000111", sales=2), source_row("Red-MR", barcode="000111", sales=3)],
            [source_row("Red", barcode="000222", sales=4), source_row("Red /OEM", barcode="000222", sales=5)],
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual((rows.iloc[0]["ยอดขาย_ASIA"], rows.iloc[0]["ยอดขาย_GREEN"]), (5, 9))
        self.assertEqual(rows.iloc[0]["หมายเหตุ"], WARNING)

    def test_different_explicit_colors_are_not_a_barcode_mismatch_match(self) -> None:
        rows = combine(
            [source_row("ก๊อกสีแดง", barcode="000111", sales=2)],
            [source_row("ก๊อกสีน้ำเงิน", barcode="000222", sales=3)],
        )

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["barcode"]), {"000111", "000222"})
        self.assertFalse(rows["barcode_mismatch"].any())
        self.assertTrue(rows["หมายเหตุ"].eq("").all())
        self.assertEqual(rows["ยอดขาย_TOTAL"].sum(), 5)

    def test_different_explicit_sizes_remain_separate(self) -> None:
        rows = combine(
            [source_row("วาล์ว 1/2", barcode="000111", sales=2)],
            [source_row("วาล์ว 3/4", barcode="000222", sales=3)],
        )

        self.assertEqual(len(rows), 2)
        self.assertFalse(rows["barcode_mismatch"].any())

    def test_multiple_same_color_barcodes_in_one_source_cannot_fan_out(self) -> None:
        rows = combine(
            [source_row("Red", barcode="000111", sales=2), source_row("Red", barcode="000333", sales=3)],
            [source_row("Red", barcode="000222", sales=4)],
        )

        self.assertEqual(len(rows), 3)
        self.assertEqual(set(rows["barcode"]), {"000111", "000222", "000333"})
        self.assertEqual(rows["ยอดขาย_TOTAL"].sum(), 9)
        self.assertFalse(rows["barcode_mismatch"].any())

    def test_known_barcode_match_does_not_make_other_same_color_barcodes_unambiguous(self) -> None:
        rows = combine(
            [source_row("Red", barcode="000111", sales=2), source_row("Red", barcode="000333", sales=3)],
            [source_row("Red", barcode="000111", sales=4), source_row("Red", barcode="000222", sales=5)],
        )

        self.assertEqual(len(rows), 3)
        known = rows.loc[rows["barcode"].eq("000111")].iloc[0]
        self.assertEqual(known["ยอดขาย_TOTAL"], 6)
        self.assertEqual(rows["ยอดขาย_TOTAL"].sum(), 14)
        self.assertFalse(rows["barcode_mismatch"].any())

    def test_distinct_color_pairs_can_each_match_their_own_changed_barcode(self) -> None:
        rows = combine(
            [source_row("Red", barcode="000111", sales=2), source_row("White", barcode="000333", sales=3)],
            [source_row("Red", barcode="000222", sales=4), source_row("White", barcode="000444", sales=5)],
        )

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["barcode"]), {"000222", "000444"})
        self.assertEqual(sorted(rows["ยอดขาย_TOTAL"].tolist()), [6, 8])
        self.assertTrue(rows["barcode_mismatch"].all())

    def test_different_suppliers_do_not_merge(self) -> None:
        asia = source_row("Red", barcode="000111", sales=2)
        green = source_row("Red", barcode="000222", sales=3)
        green["buyer"] = "A0030"
        rows = combine([asia], [green])

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["buyer"]), {"A0029", "A0030"})
        self.assertFalse(rows["barcode_mismatch"].any())

    def test_different_item_codes_with_different_barcodes_do_not_merge(self) -> None:
        rows = combine(
            [source_row("Red", code="MC-510", barcode="000111", sales=2)],
            [source_row("Red", code="MC-511", barcode="000222", sales=3)],
        )

        self.assertEqual(len(rows), 2)
        self.assertFalse(rows["barcode_mismatch"].any())

    def test_same_barcodes_and_missing_one_source_barcode_do_not_warn(self) -> None:
        for asia_barcode, green_barcode, selected in (
            ("000123", "000123", "000123"), ("", "000123", "000123"),
            ("000123", "", "000123"), ("", "", ""),
        ):
            with self.subTest(asia=asia_barcode, green=green_barcode):
                rows = combine(
                    [source_row("Red", barcode=asia_barcode, sales=2)],
                    [source_row("Red", barcode=green_barcode, sales=3)],
                )
                self.assertEqual(len(rows), 1)
                self.assertEqual(rows.iloc[0]["barcode"], selected)
                self.assertFalse(rows.iloc[0]["barcode_mismatch"])
                self.assertEqual(rows.iloc[0]["หมายเหตุ"], "")

    def test_po_uses_green_catalog_metadata_and_yellow_final_barcode(self) -> None:
        rows = self.matching_mismatch()
        with tempfile.TemporaryDirectory() as tmp:
            po = self.generate_po(rows, Path(tmp), [
                catalog_entry(barcode="000111", brand="ASIA catalog brand", carton=24),
                catalog_entry(),
            ])

            self.assertEqual(po["D9"].value, WARNING)
            self.assertEqual(po["Y9"].value, "000222")
            self.assertEqual(po["Y9"].data_type, "s")
            self.assertEqual(po["Y9"].number_format, "@")
            self.assertEqual(po["Y9"].fill.fgColor.rgb[-6:], "FFFF00")
            self.assertEqual(po["D9"].fill.fgColor.rgb[-6:], "FFFF00")
            self.assertEqual(po["E9"].value, "GREEN catalog brand")
            self.assertEqual(po["H9"].value, 12)

    def test_po_retains_mismatch_warning_with_missing_carton_note(self) -> None:
        filename = "รายการสินค้ารุ่นใหม่.xlsx"
        with tempfile.TemporaryDirectory() as tmp:
            po = self.generate_po(self.matching_mismatch(), Path(tmp), [catalog_entry(carton=None)], filename)

            self.assertEqual(po["D9"].value.splitlines(), [WARNING, f'ไม่มี "carton" ใน "{filename}"'])
            self.assertIsNone(po["H9"].value)
            self.assertEqual(po["Y9"].value, "000222")
            self.assertEqual(po["Y9"].fill.fgColor.rgb[-6:], "FFFF00")
            self.assertEqual(po["D9"].value.count(WARNING), 1)

    def test_populated_catalog_barcode_conflicts_still_raise(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, r"No catalog BARCODE match.*000222"):
                self.generate_po(self.matching_mismatch(), Path(tmp), [catalog_entry(barcode="000111")])

    def test_all_items_notes_are_adjacent_to_name_and_warning_barcode_stays_yellow(self) -> None:
        rows = self.matching_mismatch()
        rows.loc[0, "หมายเหตุ"] += "\nExisting note"
        with tempfile.TemporaryDirectory() as tmp:
            path = main.export_vendor_all_items_excel(rows, "A0029", out_folder=tmp)
            ws = openpyxl.load_workbook(path)["all_items"]
            headers = {ws.cell(1, col).value: col for col in range(1, ws.max_column + 1)}
            self.assertEqual(headers["หมายเหตุ"], headers["รายละเอียดสินค้า"] + 1)
            self.assertEqual(ws.cell(2, headers["หมายเหตุ"]).value, WARNING + "\nExisting note")
            self.assertEqual(ws.cell(2, headers["barcode"]).value, "000222")
            self.assertEqual(ws.cell(2, headers["barcode"]).fill.fgColor.rgb[-6:], "FFFF00")
            for name, expected in (("barcode", "000222"), ("barcode_ASIA", "000111"), ("barcode_GREEN", "000222")):
                cell = ws.cell(2, headers[name])
                self.assertEqual(cell.value, expected)
                self.assertEqual(cell.data_type, "s")
                self.assertEqual(cell.number_format, "@")

    def test_asia_only_barcode_is_written_without_warning(self) -> None:
        rows = combine([source_row("Red", barcode="000111", sales=6)], [])
        with tempfile.TemporaryDirectory() as tmp:
            po = self.generate_po(rows, Path(tmp), [catalog_entry(barcode="000111")])

            self.assertEqual(po["Y9"].value, "000111")
            self.assertIn(po["D9"].value, (None, ""))


if __name__ == "__main__":
    unittest.main()
