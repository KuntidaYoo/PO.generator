"""Regression checks for barcode-based PO variants and downloads."""

from __future__ import annotations

import datetime
import re
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pandas as pd
from PIL import Image as PILImage
from openpyxl.drawing.image import Image as XLImage

import main


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "ตัวอย่างใบสั่งซื้อต่างประเทศ.xlsx"


def source_row(
    description: str,
    *,
    barcode: str = "",
    sales: float = 0,
    stock: float = 0,
    on_order: float = 0,
    yuan: float = 0.447,
    code: str = "MC-510",
) -> dict:
    return {
        "buyer": "A0029",
        "barcode": barcode,
        "รหัสสินค้า": code,
        "รายละเอียดสินค้า": description,
        "ยอดขาย": sales,
        "สินค้าคงเหลือ": stock,
        "ON_ORDER": on_order,
        "หยวน": yuan,
    }


def combine(asia: list[dict], green: list[dict], months: int = 3) -> pd.DataFrame:
    return main.build_combined_all(
        pd.DataFrame(asia),
        pd.DataFrame(green),
        months=months,
        min_factor=4,
        max_factor=7,
    )


def parsed_sample_row(barcode: str, description: str) -> dict:
    """Run one screenshot-like Express line through the public parser."""
    line = f"A0029 {barcode} 01-22-1085-E1-1 MC-510 {description} 0.00 0.00 0.00 0.00 0.00 0.447"
    with (
        patch.object(main, "extract_5_6_block_from_money_chunk", return_value=[0.0] * 6),
        patch.object(main, "extract_yuan_after_money_block", return_value=0.447),
    ):
        fields = main.parse_line_to_fields(pd.Series(dtype=object), line)
    assert fields is not None
    code, clean_description = main.split_product_field(fields["สินค้า"])
    return source_row(clean_description, barcode=fields["barcode"], code=code)


def make_catalog(
    path: Path,
    carton_quantities: tuple[int | None, int | None],
    images: bool,
    descriptions: tuple[str, str] = (
        "unrelated catalog option one",
        "unrelated catalog option two",
    ),
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "A0029"
    ws.append(["item", "picture", "description", "brand", "material", "weight", "qty/carton"])
    ws.append(["MC-510", None, descriptions[0], "Brand", "Steel", 1, carton_quantities[0]])
    ws.append(["MC-510", None, descriptions[1], "Brand", "Steel", 1, carton_quantities[1]])
    if images:
        for row, color in ((2, "red"), (3, "blue")):
            png = path.parent / f"catalog_{row}.png"
            PILImage.new("RGB", (12, 12), color).save(png)
            ws.add_image(XLImage(str(png)), f"B{row}")
    wb.save(path)


def make_barcode_catalog(path: Path, entries: list[dict]) -> None:
    """Create a catalog using the new I/BARCODE column and optional pictures."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "A0029"
    ws.append([
        "BUYER ITEM NO.", "GOODS PICTURE", "GOODS DESCRIPTION", "BRAND",
        "MATERIAL", "Weight", "QTY PER CARTON", "UNIT PRICE (FOB)", "BARCODE",
    ])
    for row_number, entry in enumerate(entries, start=2):
        ws.append([
            entry.get("code", "MC-510"), None, entry["description"], entry.get("brand", "Brand"),
            entry.get("material", "Steel"), entry.get("weight", 1),
            entry.get("carton", 10), None, entry.get("barcode", ""),
        ])
        ws.cell(row_number, 9).number_format = "@"
        color = entry.get("picture_color")
        if color:
            png = path.parent / f"catalog_barcode_{row_number}.png"
            PILImage.new("RGB", (12, 12), color).save(png)
            ws.add_image(XLImage(str(png)), f"B{row_number}")
    wb.save(path)


def item_picture_colors(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[int, tuple[int, int, int]]:
    return {
        image.anchor._from.row + 1: PILImage.open(BytesIO(image._data())).convert("RGB").getpixel((0, 0))
        for image in ws._images
        if getattr(getattr(image.anchor, "_from", None), "col", None) == 1
        and getattr(getattr(image.anchor, "_from", None), "row", None) is not None
        and image.anchor._from.row >= 8
    }


class POVariantTests(unittest.TestCase):
    def test_mc510_sample_colors_keep_three_rows(self) -> None:
        # These values are the three A0029 / MC-510 entries in the PO-G sample.
        green = [
            parsed_sample_row("8858778318375", "ก็อกแฟนซีสี-แดง(R)"),
            parsed_sample_row("8858778318498", "ก็อกแฟนซีสี-เขียว(G)"),
            parsed_sample_row("8858778318504", "ก็อกแฟนซีสี-น้ำเงิน(LB)"),
        ]

        rows = combine([], green)

        self.assertEqual(len(rows), 3)
        self.assertEqual(
            set(zip(rows["barcode"], rows["รายละเอียดสินค้า"])),
            {
                ("8858778318375", "ก็อกแฟนซีสี-แดง(R)"),
                ("8858778318498", "ก็อกแฟนซีสี-เขียว(G)"),
                ("8858778318504", "ก็อกแฟนซีสี-น้ำเงิน(LB)"),
            },
        )
        self.assertTrue((rows["ยอดขาย_TOTAL"] == 0).all())
        self.assertTrue((rows["MIN_NUM"] == 0).all())
        self.assertFalse((rows["TOTAL_QTY_NUM"] < rows["MIN_NUM"]).any())

        with tempfile.TemporaryDirectory() as tmp:
            path = main.export_vendor_all_items_excel(rows, "A0029", out_folder=tmp)
            ws = openpyxl.load_workbook(path)["all_items"]
            self.assertEqual(ws.max_row, 4)
            headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
            exported = {
                (ws.cell(r, headers["barcode"]).value, ws.cell(r, headers["รายละเอียดสินค้า"]).value)
                for r in range(2, 5)
            }
            self.assertEqual(exported, set(zip(rows["barcode"], rows["รายละเอียดสินค้า"])))

    def test_matching_barcode_combines_descriptions_and_same_source_missing_barcode(self) -> None:
        red = "ก็อกแฟนซีสี-แดง(R)"
        blue = "ก็อกแฟนซีสี-น้ำเงิน(LB)"
        asia = [
            source_row(red, barcode="000123", sales=4, stock=1, on_order=1, yuan=10),
            source_row(red, barcode="", sales=3, stock=0, yuan=10),
        ]
        green = [
            source_row(red, barcode="000123", sales=2, stock=3, yuan=11),
            source_row(blue, barcode="000123", sales=12, stock=20, yuan=11),
            source_row(red, barcode="000123", sales=5, stock=0, yuan=11),
        ]

        rows = combine(asia, green)

        self.assertEqual(len(rows), 1)
        matching_red = rows[
            (rows["barcode"] == "000123")
            & (rows["รายละเอียดสินค้า"] == red)
            & (rows["หยวน"] == 11)
        ].iloc[0]
        self.assertEqual(matching_red["ยอดขาย_ASIA"], 7)
        self.assertEqual(matching_red["ยอดขาย_GREEN"], 19)
        self.assertEqual(matching_red["STOCK_ASIA"], 1)
        self.assertEqual(matching_red["STOCK_GREEN"], 23)
        self.assertEqual(matching_red["ON_ORDER_TOTAL"], 1)
        self.assertEqual(matching_red["USE_MONTH"], 9)
        self.assertEqual(matching_red["TOTAL_QTY_NUM"], 25)
        self.assertEqual(matching_red["MIN_NUM"], 36)
        self.assertEqual(matching_red["MAX_NUM"], 63)

    def test_asia_only_row_has_no_po_barcode_and_parser_keeps_leading_zero(self) -> None:
        parsed_green = parsed_sample_row("000123", "ก็อกแฟนซีสี-แดง(R)")
        self.assertEqual(parsed_green["barcode"], "000123")

        rows = combine(
            [source_row("ASIA only", barcode="999888", code="MC-512", sales=6)],
            [],
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows.iloc[0]["barcode"], "")
        self.assertEqual(rows.iloc[0]["ยอดขาย_ASIA"], 6)
        self.assertEqual(rows.iloc[0]["ยอดขาย_GREEN"], 0)

    def test_same_barcode_source_records_aggregate_with_first_description(self) -> None:
        green = [
            source_row("Red", barcode="0001", sales=2, stock=1, yuan=10),
            source_row("Red-MR", barcode="0001", sales=3, stock=4, yuan=10),
            source_row("Red", barcode="0002", sales=5, stock=6, yuan=10),
            source_row("Red-IR", barcode="0001", sales=7, stock=8, yuan=10),
        ]

        rows = combine([], green)

        self.assertEqual(len(rows), 2)
        identical = rows[(rows["barcode"] == "0001")].iloc[0]
        self.assertEqual(identical["รายละเอียดสินค้า"], "Red")
        self.assertEqual(identical["ยอดขาย_GREEN"], 12)
        self.assertEqual(identical["STOCK_GREEN"], 13)
        self.assertEqual(rows[rows["barcode"] == "0002"].iloc[0]["ยอดขาย_GREEN"], 5)

    def test_conflicting_prices_for_same_barcode_raise_clear_error(self) -> None:
        with self.assertRaisesRegex(ValueError, r"(?i)conflicting GREEN prices.*0001"):
            combine([], [
                source_row("Red", barcode="0001", sales=2, yuan=10),
                source_row("Red-MR", barcode="0001", sales=3, yuan=11),
            ])

    def test_inactive_duplicate_price_does_not_override_active_variant(self) -> None:
        rows = combine([], [
            source_row("ก๊อกสนามคอยาววาล์วเซรามิค 1/2", barcode="8858778308567",
                       code="MC401-2L", sales=1647, yuan=8.04),
            source_row("ก็อกสนามคอยาววาล์วเซรามิค B", barcode="8858778308567",
                       code="MC401-2L", yuan=6.19),
        ])
        self.assertEqual(len(rows), 1)
        row = rows.iloc[0]
        self.assertEqual(row["หยวน"], 8.04)
        self.assertEqual(row["ยอดขาย_GREEN"], 1647)
        self.assertEqual(row["รายละเอียดสินค้า"], "ก๊อกสนามคอยาววาล์วเซรามิค 1/2")

    def test_conflicting_inactive_prices_leave_price_blank(self) -> None:
        rows = combine([], [
            source_row("Same", barcode="0001", yuan=10),
            source_row("Same-MR", barcode="0001", yuan=11),
        ])
        self.assertEqual(len(rows), 1)
        self.assertTrue(pd.isna(rows.iloc[0]["หยวน"]))

    def test_1g022_suffixes_and_blank_asia_prices_use_green_identity(self) -> None:
        green_names = {
            "A-1701C": "ก๊อกอ่างล้างหน้าแนวตั้งด้ามยก",
            "A-1704K": "ก๊อกอ่างล้างหน้าด้ามยก-สีเทา",
        }
        barcodes = {"A-1701C": "0001701", "A-1704K": "0001704"}
        prices = {"A-1701C": 29.69, "A-1704K": 33.81}
        asia = [
            source_row(green_names[code] + suffix, code=code,
                       barcode=barcodes[code], sales=12, stock=2, yuan=float("nan"))
            for code, suffix in (("A-1701C", "-MR"), ("A-1704K", "-IR"))
        ]
        green = [
            source_row(name, code=code, barcode=barcodes[code],
                       sales=15, stock=3, yuan=prices[code])
            for code, name in green_names.items()
        ]
        rows = combine(asia, green)
        self.assertEqual(len(rows), 2)
        self.assertTrue((rows["TOTAL_QTY_NUM"] < rows["MIN_NUM"]).all())
        for code, name in green_names.items():
            row = rows[rows["รหัสสินค้า"] == code].iloc[0]
            self.assertEqual(row["รายละเอียดสินค้า"], name)
            self.assertEqual(row["ยอดขาย_ASIA"], 12)
            self.assertEqual(row["ยอดขาย_GREEN"], 15)
            self.assertEqual(row["หยวน"], prices[code])

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog.xlsx"
            make_barcode_catalog(catalog_path, [
                {"code": code, "barcode": barcodes[code],
                 "description": name + " catalog wording", "carton": 48}
                for code, name in green_names.items()
            ])
            all_path = main.export_vendor_all_items_excel(rows, "A0029", out_folder=tmp)
            all_ws = openpyxl.load_workbook(all_path)["all_items"]
            self.assertEqual(all_ws.max_row, 3)
            with patch.object(main, "PO_OUTPUT_FOLDER", tmp):
                po_path = main.generate_po_from_combined(
                    rows, "A0029", datetime.date(2026, 9, 27), 6,
                    str(TEMPLATE), str(catalog_path),
                    str(output / "missing_vendors.xlsx"), 4, 7,
                )
            po = openpyxl.load_workbook(po_path)["PO"]
            for row_number in (9, 10):
                code = po[f"A{row_number}"].value
                self.assertEqual(po[f"C{row_number}"].value, green_names[code])
                self.assertEqual(po[f"X{row_number}"].value, barcodes[code])
                self.assertEqual(po[f"L{row_number}"].value, prices[code])
                self.assertEqual(po[f"G{row_number}"].value, 48)

    def test_blank_barcodes_match_after_ignoring_source_marker(self) -> None:
        rows = combine(
            [source_row("Red-MR", barcode="", sales=2)],
            [source_row("Red", barcode="", sales=3)],
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows.iloc[0]["รายละเอียดสินค้า"], "Red")
        self.assertEqual(rows.iloc[0]["ยอดขาย_TOTAL"], 5)
        self.assertTrue((rows["barcode"] == "").all())

    def test_blank_barcodes_preserve_different_color_descriptions(self) -> None:
        rows = combine(
            [source_row("CBS-320 สีแดง", barcode="", code="CBS-320", sales=2)],
            [source_row("CBS-320 สีเขียว", barcode="", code="CBS-320", sales=3)],
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), {"CBS-320 สีแดง", "CBS-320 สีเขียว"})

    def test_missing_asia_barcode_merges_unique_green_name(self) -> None:
        name = "ก๊อกอ่างล้างหน้าแนวตั้งด้ามยก"
        rows = combine(
            [source_row(name + "-MR", barcode="", code="A-1701C", sales=9, stock=2,
                        yuan=float("nan"))],
            [source_row(name, barcode="0001701", code="A-1701C", sales=12, stock=3,
                        yuan=29.69)],
        )
        self.assertEqual(len(rows), 1)
        row = rows.iloc[0]
        self.assertEqual(row["barcode"], "0001701")
        self.assertEqual(row["catalog_match_barcode"], "0001701")
        self.assertEqual(row["รายละเอียดสินค้า"], name)
        self.assertEqual((row["ยอดขาย_ASIA"], row["ยอดขาย_GREEN"]), (9, 12))
        self.assertEqual((row["STOCK_ASIA"], row["STOCK_GREEN"]), (2, 3))
        self.assertEqual(row["หยวน"], 29.69)

    def test_missing_green_barcode_uses_asia_barcode_only_for_catalog(self) -> None:
        name = "ก๊อกอ่างล้างหน้าด้ามยก-สีเทา"
        rows = combine(
            [source_row(name + "-IR", barcode="0001704", code="A-1704K", sales=2)],
            [source_row(name, barcode="", code="A-1704-K", sales=3)],
        )
        self.assertEqual(len(rows), 1)
        row = rows.iloc[0]
        self.assertEqual(row["รหัสสินค้า"], "A-1704-K")
        self.assertEqual(row["barcode"], "")
        self.assertEqual(row["catalog_match_barcode"], "0001704")
        self.assertEqual(row["ยอดขาย_TOTAL"], 5)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog.xlsx"
            make_barcode_catalog(catalog_path, [
                {"code": "A-1704-K", "barcode": "0001704", "description": name,
                 "carton": 48},
            ])
            with patch.object(main, "PO_OUTPUT_FOLDER", tmp):
                po_path = main.generate_po_from_combined(
                    rows, "A0029", datetime.date(2026, 9, 27), 6,
                    str(TEMPLATE), str(catalog_path),
                    str(output / "missing_vendors.xlsx"), 4, 7,
                )
            po = openpyxl.load_workbook(po_path)["PO"]
            self.assertEqual(po["A9"].value, "A-1704-K")
            self.assertEqual(po["C9"].value, name)
            self.assertEqual(po["G9"].value, 48)
            self.assertIn(po["X9"].value, (None, ""))

    def test_item_code_parenthetical_color_is_part_of_variant_identity(self) -> None:
        rows = combine(
            [source_row("CBS-320 สีแดง", barcode="", code="CBS-320(R)", sales=2)],
            [source_row("CBS-320 สีเขียว", barcode="", code="CBS-320(G)", sales=3)],
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รหัสสินค้า"]), {"CBS-320(R)", "CBS-320(G)"})
        self.assertEqual(set(rows["ยอดขาย_TOTAL"]), {2, 3})

    def test_blank_barcode_cannot_fan_out_to_multiple_color_variants(self) -> None:
        name = "CBS-320 same wording"
        rows = combine(
            [source_row(name, barcode="", code="CBS-320", sales=2)],
            [source_row(name, barcode="0001", code="CBS-320", sales=3),
             source_row(name, barcode="0002", code="CBS-320", sales=4)],
        )
        self.assertEqual(len(rows), 3)
        self.assertEqual(sorted(rows["ยอดขาย_TOTAL"].tolist()), [2, 3, 4])

    def test_different_populated_barcodes_never_merge_by_name(self) -> None:
        name = "CBS-320 same wording"
        rows = combine(
            [source_row(name, barcode="0001", code="CBS-320", sales=2)],
            [source_row(name, barcode="0002", code="CBS-320", sales=3)],
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(sorted(rows["ยอดขาย_TOTAL"].tolist()), [2, 3])

    def test_blank_barcode_variant_keeps_legacy_first_price(self) -> None:
        rows = combine([], [
            source_row("Gray-IR", barcode="", code="A-1704K", sales=2, yuan=10),
            source_row("Gray", barcode="", code="A-1704-K", sales=3, yuan=11),
        ])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows.iloc[0]["หยวน"], 10)
        self.assertEqual(rows.iloc[0]["ยอดขาย_TOTAL"], 5)

    def test_parser_keeps_thai_parenthetical_color_in_code(self) -> None:
        parsed = parsed_sample_row("000123", "ก๊อกสีเขียว")
        self.assertEqual(parsed["รหัสสินค้า"], "MC-510")
        code, description = main.split_product_field(
            "01-22-1085-E1-1 MC-501(เขียว) ก๊อกแฟนซีสีเขียว"
        )
        self.assertEqual(code, "MC-501(เขียว)")
        self.assertEqual(description, "ก๊อกแฟนซีสีเขียว")

        code, description = main.split_product_field(
            "01-22-1085-E1-1 A-75-2(ชุดสายฉีดชำระ-ขา DM-948) สายฉีดชำระ"
        )
        self.assertEqual(code, "A-75-2(ชุดสายฉีดชำระ-ขา DM-948)")
        self.assertEqual(description, "สายฉีดชำระ")
        code, description = main.split_product_field(
            "01-25-2117 SL-3305C(H)-CRหัวฝักบัวสีแดง"
        )
        self.assertEqual(code, "SL-3305C(H)-CR")
        self.assertEqual(description, "หัวฝักบัวสีแดง")

    def test_trailing_period_barcode_and_parenthetical_code_from_express(self) -> None:
        line = (
            "1G022 8858778325991. } 01-25-2117 BM-120(WS) "
            "สายอ่อนฝักบัวเยอรมัน-สีดำ 120CM. 2.00 71.20 2.00 71.20 0.00"
        )
        with (
            patch.object(main, "extract_5_6_block_from_money_chunk", return_value=[0.0] * 6),
            patch.object(main, "extract_yuan_after_money_block", return_value=1.0),
        ):
            fields = main.parse_line_to_fields(pd.Series(dtype=object), line)
        self.assertIsNotNone(fields)
        self.assertEqual(fields["barcode"], "8858778325991")
        code, description = main.split_product_field(fields["สินค้า"])
        self.assertEqual(code, "BM-120(WS)")
        self.assertTrue(description.startswith("สายอ่อนฝักบัว"))
        self.assertEqual(main.normalize_barcode("8858778325991."), "8858778325991")
        self.assertEqual(main.normalize_barcode("1234567."), "1234567.")
        code, description = main.split_product_field(
            "01-22-1085-E1-1 MC-501(เขียว)ก๊อกแฟนซีสีเขียว"
        )
        self.assertEqual(code, "MC-501(เขียว)")
        self.assertEqual(description, "ก๊อกแฟนซีสีเขียว")

    def test_missing_barcode_in_same_source_attaches_to_unique_variant(self) -> None:
        rows = combine([], [
            source_row("CBS-320 สีแดง", barcode="0001", code="CBS-320", sales=2,
                       stock=1, yuan=10),
            source_row("CBS-320 สีแดง-MR", barcode="", code="CBS-320", sales=3,
                       stock=4, yuan=10),
        ])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows.iloc[0]["barcode"], "0001")
        self.assertEqual(rows.iloc[0]["ยอดขาย_GREEN"], 5)
        self.assertEqual(rows.iloc[0]["STOCK_GREEN"], 5)
        self.assertEqual(rows.iloc[0]["รายละเอียดสินค้า"], "CBS-320 สีแดง")

        reverse_rows = combine([], [
            source_row("CBS-320 สีแดง-MR", barcode="", code="CBS-320", sales=3,
                       stock=4, yuan=10),
            source_row("CBS-320 สีแดง", barcode="0001", code="CBS-320", sales=2,
                       stock=1, yuan=10),
        ])
        self.assertEqual(len(reverse_rows), 1)
        self.assertEqual(reverse_rows.iloc[0]["รายละเอียดสินค้า"], "CBS-320 สีแดง")
        self.assertEqual(reverse_rows.iloc[0]["ยอดขาย_GREEN"], 5)

    def test_unique_barcode_matches_across_different_express_codes(self) -> None:
        rows = combine(
            [source_row("ASIA set description", barcode="000155", code="A-155(ชุด)",
                        sales=2)],
            [source_row("GREEN product wording", barcode="000155", code="A-155",
                        sales=3)],
        )
        self.assertEqual(len(rows), 1)
        row = rows.iloc[0]
        self.assertEqual(row["รหัสสินค้า"], "A-155")
        self.assertEqual(row["รายละเอียดสินค้า"], "GREEN product wording")
        self.assertEqual(row["barcode"], "000155")
        self.assertEqual(row["ยอดขาย_TOTAL"], 5)

        color_code_rows = combine(
            [source_row("Tap red", barcode="000320", code="CBS-320(R)", sales=2)],
            [source_row("Tap red", barcode="000320", code="CBS-320", sales=3)],
        )
        self.assertEqual(len(color_code_rows), 1)
        self.assertEqual(color_code_rows.iloc[0]["ยอดขาย_TOTAL"], 5)

        set_rows = combine(
            [source_row("Set description", barcode="2024113600127",
                        code="A-155(ชุดท่อน้ำทิ้ง+สะดือ8นิ้ว)", sales=3741)],
            [source_row("Green wording", barcode="2024113600127", code="A-155",
                        sales=914)],
        )
        self.assertEqual(len(set_rows), 1)
        self.assertEqual(set_rows.iloc[0]["ยอดขาย_TOTAL"], 4655)

    def test_missing_barcode_in_same_source_does_not_fan_out(self) -> None:
        rows = combine([], [
            source_row("CBS-320 สีแดง", barcode="0001", code="CBS-320", sales=2),
            source_row("CBS-320 สีแดง", barcode="0002", code="CBS-320", sales=3),
            source_row("CBS-320 สีแดง-MR", barcode="", code="CBS-320", sales=4),
        ])
        self.assertEqual(len(rows), 3)
        self.assertEqual(sorted(rows["ยอดขาย_GREEN"].tolist()), [2, 3, 4])
        self.assertEqual(sorted(rows["barcode"].tolist()), ["", "0001", "0002"])

    def test_missing_normalization_values_are_blank(self) -> None:
        self.assertEqual(main.normalize_item_code(float("nan")), "")
        self.assertEqual(main.normalize_product_description(float("nan")), "")

    def test_totals_extend_past_template_item_rows(self) -> None:
        rows = combine([], [
            source_row(f"Variant {index}", barcode=f"000{index}", sales=3)
            for index in range(1, 8)
        ])
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog.xlsx"
            make_barcode_catalog(catalog_path, [
                {"barcode": f"000{index}", "description": f"Variant {index}", "carton": 10}
                for index in range(1, 8)
            ])
            with patch.object(main, "PO_OUTPUT_FOLDER", tmp):
                po_path = main.generate_po_from_combined(
                    rows, "A0029", datetime.date(2026, 9, 27), 6,
                    str(TEMPLATE), str(catalog_path),
                    str(output / "missing_vendors.xlsx"), 4, 7,
                )
            po = openpyxl.load_workbook(po_path)["PO"]
            for column in ("H", "K", "N"):
                self.assertEqual(po[f"{column}16"].value, f"=SUM({column}9:{column}15)")
            self.assertEqual(po["N17"].value, 6)
            self.assertEqual(po["N18"].value, "=N16*N17")

    def test_missing_carton_quantity_raises_before_writing_division_formula(self) -> None:
        rows = combine([], [source_row("Red", barcode="000123", sales=3)])
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog.xlsx"
            make_barcode_catalog(catalog_path, [
                {"barcode": "000123", "description": "Red", "carton": None},
            ])
            with patch.object(main, "PO_OUTPUT_FOLDER", tmp):
                with self.assertRaisesRegex(ValueError, r"QTY PER CARTON.*000123"):
                    main.generate_po_from_combined(
                        rows, "A0029", datetime.date(2026, 9, 27), 6,
                        str(TEMPLATE), str(catalog_path),
                        str(output / "missing_vendors.xlsx"), 4, 7,
                    )

    def test_all_items_and_po_keep_leading_zero_barcode_as_text(self) -> None:
        description = "ก็อกแฟนซีสี-แดง(R)"
        rows = combine(
            [source_row(description, barcode="000123", sales=4, stock=1, on_order=1, yuan=10)],
            [source_row(description, barcode="000123", sales=2, stock=3, yuan=10)],
        )
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog.xlsx"
            make_barcode_catalog(catalog_path, [
                {"barcode": "000123", "description": description, "carton": 10},
            ])
            all_items_path = main.export_vendor_all_items_excel(rows, "A0029", out_folder=str(output))
            all_items = openpyxl.load_workbook(all_items_path)["all_items"]
            headers = {all_items.cell(1, col).value: col for col in range(1, all_items.max_column + 1)}
            all_barcode = all_items.cell(2, headers["barcode"])
            self.assertEqual(all_barcode.value, "000123")
            self.assertEqual(all_barcode.data_type, "s")

            with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
                po_path = main.generate_po_from_combined(
                    rows,
                    "A0029",
                    datetime.date(2026, 9, 24),
                    6,
                    str(TEMPLATE),
                    str(catalog_path),
                    str(output / "missing_vendors.xlsx"),
                    4,
                    7,
                )
            po = openpyxl.load_workbook(po_path)["PO"]
            self.assertEqual(po["X8"].value, "BARCODE")
            self.assertEqual(po["X9"].value, "000123")
            self.assertEqual(po["X9"].data_type, "s")
            self.assertEqual(po["C9"].value, description)
            self.assertEqual(po["R9"].value, 3)
            self.assertEqual(po["S9"].value, 1)
            self.assertEqual(po["T9"].value, 1)
            self.assertEqual(po["O9"].value, 2)
            self.assertIn("$X", str(po.print_area))

    def test_two_po_variants_keep_independent_prices_and_formula_rows(self) -> None:
        red = "ก็อกแฟนซีสี-แดง(R)"
        blue = "ก็อกแฟนซีสี-น้ำเงิน(LB)"
        rows = combine(
            [],
            [
                source_row(red, barcode="000123", sales=6, yuan=10.5),
                source_row(blue, barcode="000456", sales=9, yuan=12.25),
            ],
        )
        self.assertTrue((rows["TOTAL_QTY_NUM"] < rows["MIN_NUM"]).all())

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog.xlsx"
            make_catalog(catalog_path, (10, 10), images=False, descriptions=(red, blue))
            with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
                po_path = main.generate_po_from_combined(
                    rows,
                    "A0029",
                    datetime.date(2026, 9, 24),
                    6,
                    str(TEMPLATE),
                    str(catalog_path),
                    str(output / "missing_vendors.xlsx"),
                    4,
                    7,
                )
            po = openpyxl.load_workbook(po_path)["PO"]
            by_barcode = {po[f"X{r}"].value: r for r in (9, 10)}
            self.assertEqual(set(by_barcode), {"000123", "000456"})

            for barcode, description, price in (
                ("000123", red, 10.5),
                ("000456", blue, 12.25),
            ):
                row = by_barcode[barcode]
                self.assertEqual(po[f"C{row}"].value, description)
                self.assertEqual(po[f"G{row}"].value, 10)
                self.assertEqual(po[f"L{row}"].value, price)
                self.assertAlmostEqual(po[f"M{row}"].value, price * 6)
                self.assertEqual(po[f"N{row}"].value, f"=L{row}*K{row}")
                self.assertEqual(po[f"K{row}"].value, f"=I{row}+J{row}")
                self.assertEqual(po[f"X{row}"].data_type, "s")

            total_formula = po["N14"].value
            total_range = re.fullmatch(r"=SUM\(N(\d+):N(\d+)\)", total_formula or "")
            self.assertIsNotNone(total_range)
            self.assertLessEqual(int(total_range.group(1)), 9)
            self.assertGreaterEqual(int(total_range.group(2)), 10)
            self.assertIn("$X", str(po.print_area))

    def test_exact_catalog_descriptions_select_matching_images(self) -> None:
        red = "ก็อกแฟนซีสี-แดง(R)"
        blue = "ก็อกแฟนซีสี-น้ำเงิน(LB)"
        rows = combine(
            [],
            [
                source_row(red, barcode="000123", sales=3),
                source_row(blue, barcode="000456", sales=3),
            ],
        )

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog.xlsx"
            make_catalog(catalog_path, (10, 10), images=True, descriptions=(red, blue))
            with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
                po_path = main.generate_po_from_combined(
                    rows,
                    "A0029",
                    datetime.date(2026, 9, 24),
                    6,
                    str(TEMPLATE),
                    str(catalog_path),
                    str(output / "missing_vendors.xlsx"),
                    4,
                    7,
                )
            po = openpyxl.load_workbook(po_path)["PO"]
            selected = {
                image.anchor._from.row + 1: PILImage.open(BytesIO(image._data())).convert("RGB").getpixel((0, 0))
                for image in po._images
                if getattr(getattr(image.anchor, "_from", None), "col", None) == 1
                and getattr(getattr(image.anchor, "_from", None), "row", None) in (8, 9)
            }
            self.assertEqual(set(selected), {9, 10})
            for row in (9, 10):
                expected_color = (255, 0, 0) if po[f"C{row}"].value == red else (0, 0, 255)
                self.assertEqual(selected[row], expected_color)

    def test_legacy_catalog_image_uses_last_code_row(self) -> None:
        rows = combine(
            [],
            [source_row("ก็อกแฟนซีสี-แดง(R)", barcode="000123", sales=3)],
        )
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog.xlsx"
            make_catalog(catalog_path, (10, 10), images=True)
            with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
                po_path = main.generate_po_from_combined(
                    rows,
                    "A0029",
                    datetime.date(2026, 9, 24),
                    6,
                    str(TEMPLATE),
                    str(catalog_path),
                    str(output / "missing_vendors.xlsx"),
                    4,
                    7,
                    variant_counts_by_code={"MC-510": 2},
                )
            po = openpyxl.load_workbook(po_path)["PO"]
            item_images = [
                image for image in po._images
                if getattr(getattr(image.anchor, "_from", None), "row", None) == 8
                and getattr(getattr(image.anchor, "_from", None), "col", None) == 1
            ]
            self.assertEqual(len(item_images), 1)
            color = PILImage.open(BytesIO(item_images[0]._data())).convert("RGB").getpixel((0, 0))
            self.assertEqual(color, (0, 0, 255))
            self.assertEqual(po["G9"].value, 10)

    def test_legacy_catalog_uses_last_code_rows_carton_quantity(self) -> None:
        rows = combine([], [source_row("ก็อกแฟนซีสี-แดง(R)", barcode="000123", sales=3)])
        for quantities in ((10, 20), (None, 10)):
            with self.subTest(carton_quantities=quantities), tempfile.TemporaryDirectory() as tmp:
                output = Path(tmp)
                catalog_path = output / "catalog.xlsx"
                make_catalog(catalog_path, quantities, images=False)
                with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
                    po_path = main.generate_po_from_combined(
                        rows, "A0029", datetime.date(2026, 9, 24), 6,
                        str(TEMPLATE), str(catalog_path), str(output / "missing_vendors.xlsx"),
                        4, 7, variant_counts_by_code={"MC-510": 2},
                    )
                po = openpyxl.load_workbook(po_path)["PO"]
                self.assertEqual(po["G9"].value, quantities[-1])
                self.assertEqual(po["X9"].value, "000123")

    def test_catalog_barcode_selects_color_pictures_and_metadata(self) -> None:
        # Both Express rows have the same wording; only their GREEN barcodes
        # identify the different catalog colors and their different carton sizes.
        rows = combine(
            [],
            [
                source_row("MC-510 fancy faucet", barcode="000123", sales=6),
                source_row("MC-510 fancy faucet", barcode="000456", sales=9),
            ],
        )
        self.assertEqual(len(rows), 2)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog_with_barcodes.xlsx"
            make_barcode_catalog(catalog_path, [
                {
                    "barcode": "000123", "description": "MC-510 ก๊อกแฟนซีสี-แดง(R)",
                    "brand": "Red brand", "material": "Red metal", "weight": 1.25,
                    "carton": 12, "picture_color": "red",
                },
                {
                    "barcode": "000456", "description": "MC-510 ก๊อกแฟนซีสี-น้ำเงิน(LB)",
                    "brand": "Blue brand", "material": "Blue metal", "weight": 2.5,
                    "carton": 24, "picture_color": "blue",
                },
            ])
            with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
                po_path = main.generate_po_from_combined(
                    rows, "A0029", datetime.date(2026, 9, 24), 6,
                    str(TEMPLATE), str(catalog_path), str(output / "missing_vendors.xlsx"),
                    4, 7,
                )
            po = openpyxl.load_workbook(po_path)["PO"]
            by_barcode = {po[f"X{row}"].value: row for row in (9, 10)}
            self.assertEqual(set(by_barcode), {"000123", "000456"})
            self.assertTrue(all(po[f"X{row}"].data_type == "s" for row in by_barcode.values()))
            expected = {
                "000123": ("Red brand", "Red metal", 1.25, 12, (255, 0, 0)),
                "000456": ("Blue brand", "Blue metal", 2.5, 24, (0, 0, 255)),
            }
            pictures = item_picture_colors(po)
            for barcode, (brand, material, weight, carton, color) in expected.items():
                row = by_barcode[barcode]
                self.assertEqual(po[f"C{row}"].value, "MC-510 fancy faucet")
                self.assertEqual(po[f"D{row}"].value, brand)
                self.assertEqual(po[f"E{row}"].value, material)
                self.assertEqual(po[f"F{row}"].value, weight)
                self.assertEqual(po[f"G{row}"].value, carton)
                self.assertEqual(pictures[row], color)

    def test_catalog_barcode_matches_across_item_code_typo(self) -> None:
        green_description = "ก๊อกอ่างล้างหน้าด้ามยก-สีเทา"
        rows = combine([], [source_row(
            green_description, code="A-1704K", barcode="0001704", sales=3,
        )])
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog.xlsx"
            make_barcode_catalog(catalog_path, [
                {
                    "code": "A-1704-K", "barcode": "0001704",
                    "description": "catalog wording - gray", "brand": "Donmark",
                    "material": "ZINC", "weight": 501, "carton": 48,
                    "picture_color": "red",
                },
            ])
            with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
                po_path = main.generate_po_from_combined(
                    rows, "A0029", datetime.date(2026, 9, 27), 6,
                    str(TEMPLATE), str(catalog_path),
                    str(output / "missing_vendors.xlsx"), 4, 7,
                )
            po = openpyxl.load_workbook(po_path)["PO"]
            self.assertEqual(po["A9"].value, "A-1704K")
            self.assertEqual(po["C9"].value, green_description)
            self.assertEqual(po["D9"].value, "Donmark")
            self.assertEqual(po["E9"].value, "ZINC")
            self.assertEqual(po["F9"].value, 501)
            self.assertEqual(po["G9"].value, 48)
            self.assertEqual(po["X9"].value, "0001704")
            self.assertEqual(item_picture_colors(po), {9: (255, 0, 0)})

    def test_catalog_barcode_matches_row_without_item_code(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            catalog_path = Path(tmp) / "catalog.xlsx"
            make_barcode_catalog(catalog_path, [
                {"code": None, "barcode": "0001704", "description": "Gray", "carton": 48},
            ])
            catalog = main.build_catalog_map(str(catalog_path), "A0029")
            selected = main.resolve_catalog_variant(
                catalog, "A-1704K", "Green source description", 1,
                barcode="0001704",
            )
            self.assertEqual(selected["qty_per_carton"], 48)

    def test_unmatched_catalog_barcode_cannot_copy_other_variants_fields(self) -> None:
        rows = combine([], [source_row("MC-510 ก๊อกแฟนซีสี-แดง(R)", barcode="000999", sales=3)])
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog_with_barcodes.xlsx"
            make_barcode_catalog(catalog_path, [
                {
                    "barcode": "000123", "description": "MC-510 ก๊อกแฟนซีสี-แดง(R)",
                    "carton": 10, "picture_color": "red",
                },
                {
                    "barcode": "000456", "description": "MC-510 ก๊อกแฟนซีสี-น้ำเงิน(LB)",
                    "carton": 10, "picture_color": "blue",
                },
            ])
            with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
                with self.assertRaisesRegex(ValueError, r"(?is)(?=.*BARCODE)(?=.*000999)(?=.*QTY PER CARTON)"):
                    main.generate_po_from_combined(
                        rows, "A0029", datetime.date(2026, 9, 24), 6,
                        str(TEMPLATE), str(catalog_path),
                        str(output / "missing_vendors.xlsx"), 4, 7,
                    )

    def test_unmatched_barcode_cannot_use_single_other_variants_carton_size(self) -> None:
        rows = combine([], [source_row("MC-510 unknown color", barcode="000999", sales=3)])
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "single_other_variant.xlsx"
            make_barcode_catalog(catalog_path, [
                {
                    "barcode": "000123", "description": "MC-510 ก๊อกแฟนซีสี-แดง(R)",
                    "carton": 24, "picture_color": "red",
                },
            ])
            with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
                with self.assertRaisesRegex(
                    ValueError, r"(?is)(?=.*BARCODE)(?=.*QTY PER CARTON)"
                ):
                    main.generate_po_from_combined(
                        rows, "A0029", datetime.date(2026, 9, 24), 6,
                        str(TEMPLATE), str(catalog_path), str(output / "missing_vendors.xlsx"),
                        4, 7,
                    )

    def test_blank_catalog_barcodes_still_match_legacy_description(self) -> None:
        red = "MC-510 ก๊อกแฟนซีสี-แดง(R)"
        rows = combine([], [source_row(red, barcode="000123", sales=3)])
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog_with_blank_barcodes.xlsx"
            make_barcode_catalog(catalog_path, [
                {"barcode": "", "description": red, "picture_color": "red"},
                {"barcode": "", "description": "MC-510 ก๊อกแฟนซีสี-น้ำเงิน(LB)", "picture_color": "blue"},
            ])
            with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
                po_path = main.generate_po_from_combined(
                    rows, "A0029", datetime.date(2026, 9, 24), 6,
                    str(TEMPLATE), str(catalog_path), str(output / "missing_vendors.xlsx"),
                    4, 7,
                )
            po = openpyxl.load_workbook(po_path)["PO"]
            self.assertEqual(po["C9"].value, red)
            self.assertEqual(item_picture_colors(po), {9: (255, 0, 0)})

    def test_partial_catalog_barcodes_allow_blank_row_code_fallback(self) -> None:
        red = "MC-510 ก๊อกแฟนซีสี-แดง(R)"
        blue = "MC-510 ก๊อกแฟนซีสี-น้ำเงิน(LB)"
        rows = combine([], [source_row(blue, barcode="000456", sales=3)])
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "partially_updated_catalog.xlsx"
            make_barcode_catalog(catalog_path, [
                {
                    "barcode": "000123", "description": red, "carton": 12,
                    "brand": "Red brand", "picture_color": "red",
                },
                {
                    "barcode": "", "description": blue, "carton": 24,
                    "brand": "Blue brand", "picture_color": "blue",
                },
            ])
            with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
                po_path = main.generate_po_from_combined(
                    rows, "A0029", datetime.date(2026, 9, 27), 6,
                    str(TEMPLATE), str(catalog_path),
                    str(output / "missing_vendors.xlsx"), 4, 7,
                )
            po = openpyxl.load_workbook(po_path)["PO"]
            self.assertEqual(po["A9"].value, "MC-510")
            self.assertEqual(po["C9"].value, blue)
            self.assertEqual(po["D9"].value, "Blue brand")
            self.assertEqual(po["G9"].value, 24)
            self.assertEqual(po["X9"].value, "000456")
            self.assertEqual(item_picture_colors(po), {9: (0, 0, 255)})

            catalog = main.build_catalog_map(str(catalog_path), "A0029")
            for description, count in (("unmatched source wording", 1), (blue, 2)):
                selected = main.resolve_catalog_variant(
                    catalog, "MC-510", description, count,
                    description_variant_count=2, barcode="000456",
                )
                self.assertEqual(selected["brand"], "Blue brand")
                self.assertEqual(selected["qty_per_carton"], 24)

    def test_missing_source_barcode_can_match_unique_catalog_description(self) -> None:
        red = "MC-510 ก๊อกแฟนซีสี-แดง(R)"
        rows = combine([], [source_row(red, barcode="", sales=3)])
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog_with_barcodes.xlsx"
            make_barcode_catalog(catalog_path, [
                {"barcode": "000123", "description": red, "picture_color": "red"},
                {"barcode": "000456", "description": "MC-510 ก๊อกแฟนซีสี-น้ำเงิน(LB)", "picture_color": "blue"},
            ])
            with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
                po_path = main.generate_po_from_combined(
                    rows, "A0029", datetime.date(2026, 9, 24), 6,
                    str(TEMPLATE), str(catalog_path), str(output / "missing_vendors.xlsx"),
                    4, 7,
                )
            po = openpyxl.load_workbook(po_path)["PO"]
            self.assertIsNone(po["X9"].value)
            self.assertEqual(po["C9"].value, red)
            self.assertEqual(item_picture_colors(po), {9: (255, 0, 0)})

    def test_asia_only_barcode_matches_catalog_but_po_barcode_stays_blank(self) -> None:
        rough_description = "MC-510 fancy faucet"
        catalog_description = "MC-510 ก๊อกแฟนซีสี-แดง(R)"
        rows = combine(
            [source_row(rough_description, barcode="000123", sales=3)], []
        )
        self.assertEqual(rows.iloc[0]["barcode"], "")

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog_with_barcodes.xlsx"
            make_barcode_catalog(catalog_path, [
                {
                    "barcode": "000123", "description": catalog_description,
                    "carton": 12, "picture_color": "red",
                },
                {
                    "barcode": "000456", "description": "MC-510 ก๊อกแฟนซีสี-น้ำเงิน(LB)",
                    "carton": 24, "picture_color": "blue",
                },
            ])
            with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
                po_path = main.generate_po_from_combined(
                    rows, "A0029", datetime.date(2026, 9, 24), 6,
                    str(TEMPLATE), str(catalog_path), str(output / "missing_vendors.xlsx"),
                    4, 7,
                )
            po = openpyxl.load_workbook(po_path)["PO"]
            self.assertIsNone(po["X9"].value)
            self.assertEqual(po["C9"].value, rough_description)
            self.assertEqual(po["G9"].value, 12)
            self.assertEqual(item_picture_colors(po), {9: (255, 0, 0)})

    def test_shared_source_barcode_combines_descriptions_and_uses_one_picture(self) -> None:
        red = "MC-510 ก๊อกแฟนซีสี-แดง(R)"
        blue = "MC-510 ก๊อกแฟนซีสี-น้ำเงิน(LB)"
        rows = combine([], [
            source_row(red, barcode="000123", sales=3),
            source_row(blue, barcode="000123", sales=3),
        ])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows.iloc[0]["รายละเอียดสินค้า"], red)
        self.assertEqual(rows.iloc[0]["ยอดขาย_GREEN"], 6)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog_with_barcodes.xlsx"
            make_barcode_catalog(catalog_path, [
                {"barcode": "000123", "description": red, "picture_color": "red"},
            ])
            with patch.object(main, "PO_OUTPUT_FOLDER", str(output)):
                po_path = main.generate_po_from_combined(
                    rows, "A0029", datetime.date(2026, 9, 24), 6,
                    str(TEMPLATE), str(catalog_path), str(output / "missing_vendors.xlsx"),
                    4, 7,
                )
            po = openpyxl.load_workbook(po_path)["PO"]
            self.assertEqual(po["C9"].value, red)
            self.assertEqual(po["X9"].value, "000123")
            self.assertEqual(item_picture_colors(po), {9: (255, 0, 0)})

    def test_duplicate_catalog_barcode_does_not_choose_matching_description(self) -> None:
        red = "MC-510 ก๊อกแฟนซีสี-แดง(R)"
        blue = "MC-510 ก๊อกแฟนซีสี-น้ำเงิน(LB)"
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "duplicate_barcodes.xlsx"
            make_barcode_catalog(catalog_path, [
                {"barcode": "000123", "description": red, "carton": 10, "picture_color": "red"},
                {"barcode": "000123", "description": blue, "carton": 20, "picture_color": "blue"},
            ])
            catalog_map = main.build_catalog_map(str(catalog_path), "A0029")
            with self.assertRaisesRegex((ValueError, RuntimeError), r"(?i)carton"):
                main.resolve_catalog_variant(catalog_map, "MC-510", "unknown color", 1, barcode="000123")
            with self.assertRaisesRegex(ValueError, r"(?i)QTY PER CARTON.*000123"):
                main.resolve_catalog_variant(catalog_map, "MC-510", red, 1, barcode="000123")

    def test_duplicate_catalog_barcode_across_codes_uses_shared_fields_only(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            catalog_path = Path(tmp) / "duplicate_barcodes.xlsx"
            make_barcode_catalog(catalog_path, [
                {
                    "code": "A-1704-K", "barcode": "0001704", "description": "Gray",
                    "brand": "Donmark", "material": "ZINC", "weight": 501,
                    "carton": 48, "picture_color": "red",
                },
                {
                    "code": "MISTYPED", "barcode": "0001704", "description": "Other",
                    "brand": "Donmark", "material": "BRASS", "weight": 501,
                    "carton": 48, "picture_color": "blue",
                },
            ])
            catalog = main.build_catalog_map(str(catalog_path), "A0029")
            selected = main.resolve_catalog_variant(
                catalog, "A-1704K", "Gray", 1, barcode="0001704"
            )
            self.assertEqual(selected["brand"], "Donmark")
            self.assertEqual(selected["material"], "")
            self.assertEqual(selected["weight"], 501)
            self.assertEqual(selected["qty_per_carton"], 48)
            self.assertIsNone(selected["img_bytes"])

    def test_catalog_name_fallback_handles_code_punctuation_and_source_marker(self) -> None:
        name = "ก๊อกอ่างล้างหน้าด้ามยก-สีเทา"
        with tempfile.TemporaryDirectory() as tmp:
            catalog_path = Path(tmp) / "catalog.xlsx"
            make_barcode_catalog(catalog_path, [
                {"code": "A-1704-K", "barcode": "", "description": name,
                 "carton": 48, "picture_color": "red"},
            ])
            catalog = main.build_catalog_map(str(catalog_path), "A0029")
            selected = main.resolve_catalog_variant(
                catalog, "A-1704K", name + "-IR", 1, barcode="0001704"
            )
            self.assertEqual(selected["qty_per_carton"], 48)
            self.assertIsNotNone(selected["img_bytes"])

    def test_catalog_name_fallback_when_source_barcode_missing(self) -> None:
        red = "CBS-320 สีแดง(R)"
        blue = "CBS-320 สีฟ้า(LB)"
        with tempfile.TemporaryDirectory() as tmp:
            catalog_path = Path(tmp) / "catalog.xlsx"
            make_barcode_catalog(catalog_path, [
                {"code": "CBS-320", "barcode": "0001", "description": red,
                 "carton": 12, "picture_color": "red"},
                {"code": "CBS-320", "barcode": "0002", "description": blue,
                 "carton": 24, "picture_color": "blue"},
            ])
            catalog = main.build_catalog_map(str(catalog_path), "A0029")
            selected = main.resolve_catalog_variant(
                catalog, "CBS-320", red + "-MR", 2, barcode=""
            )
            self.assertEqual(selected["qty_per_carton"], 12)
            self.assertIsNotNone(selected["img_bytes"])

    def test_unique_catalog_code_fallback_with_different_wording(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            catalog_path = Path(tmp) / "catalog.xlsx"
            make_barcode_catalog(catalog_path, [
                {"code": "A-1704-K", "barcode": "", "description": "older catalog wording",
                 "carton": 48},
            ])
            catalog = main.build_catalog_map(str(catalog_path), "A0029")
            selected = main.resolve_catalog_variant(
                catalog, "A-1704K", "ก๊อกอ่างล้างหน้าด้ามยก-สีเทา", 1,
                barcode="0001704"
            )
            self.assertEqual(selected["qty_per_carton"], 48)

    def test_catalog_conflicting_barcode_blocks_name_fallback(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            catalog_path = Path(tmp) / "catalog.xlsx"
            make_barcode_catalog(catalog_path, [
                {"code": "A-1704-K", "barcode": "0009999", "description": "Gray",
                 "carton": 48},
            ])
            catalog = main.build_catalog_map(str(catalog_path), "A0029")
            with self.assertRaisesRegex(ValueError, r"(?i)No catalog BARCODE match"):
                main.resolve_catalog_variant(
                    catalog, "A-1704K", "Gray", 1, barcode="0001704"
                )

    def test_legacy_catalog_allows_multiple_source_variants_for_one_code_row(self) -> None:
        rows = combine([], [
            source_row("Gray", code="A-1704K", barcode="0001", sales=3),
            source_row("Blue", code="A-1704-K", barcode="0002", sales=4),
        ])
        self.assertEqual(len(rows), 2)
        by_code, _, _ = main.catalog_variant_counts(rows, "catalog_match_barcode")
        self.assertEqual(by_code, {"a1704k": 2})

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "catalog.xlsx"
            make_barcode_catalog(catalog_path, [
                {"code": "A-1704-K", "barcode": "", "description": "old wording",
                 "carton": 48},
            ])
            with patch.object(main, "PO_OUTPUT_FOLDER", tmp):
                po_path = main.generate_po_from_combined(
                    rows, "A0029", datetime.date(2026, 9, 27), 6,
                    str(TEMPLATE), str(catalog_path),
                    str(output / "missing_vendors.xlsx"), 4, 7,
                )
            po = openpyxl.load_workbook(po_path)["PO"]
            self.assertEqual([po[f"G{r}"].value for r in (9, 10)], [48, 48])
            self.assertEqual({po[f"X{r}"].value for r in (9, 10)}, {"0001", "0002"})



if __name__ == "__main__":
    unittest.main()
