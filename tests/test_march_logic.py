"""March item-code behavior with an output-only barcode column."""
import datetime
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pandas as pd
import main

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "ตัวอย่างใบสั่งซื้อต่างประเทศ.xlsx"


def source(code="MC-1", barcode="", sales=6, stock=0, price=10, description="Red"):
    return {"buyer": "1G022", "รหัสสินค้า": code, "barcode": barcode,
            "รายละเอียดสินค้า": description, "ยอดขาย": sales,
            "สินค้าคงเหลือ": stock, "ON_ORDER": 0, "หยวน": price}


def combine(asia, green):
    return main.build_combined_all(pd.DataFrame(asia), pd.DataFrame(green), 3, 4, 7)


class MarchLogicTests(unittest.TestCase):
    def test_same_code_groups_colors_barcodes_and_keeps_first_source_price(self):
        result = combine([source(barcode="001", sales=6, price=5)], [
            source(barcode="002", sales=9, stock=2, price=10, description="Red"),
            source(barcode="003", sales=3, stock=1, price=20, description="Blue"),
        ])
        self.assertEqual(len(result), 1)
        row = result.iloc[0]
        self.assertEqual(row["ยอดขาย_TOTAL"], 18)
        self.assertEqual(row["STOCK_GREEN"], 3)
        self.assertEqual(row["หยวน"], 10)
        self.assertEqual(row["barcode"], "002")
        self.assertEqual(row["USE_MONTH"], 6)
        self.assertEqual(row["MIN_NUM"], 24)
        self.assertEqual(row["MAX_NUM"], 42)
        self.assertNotIn("หมายเหตุ", result.columns)

    def test_exact_item_codes_remain_separate_even_with_same_barcode(self):
        rows = combine([source(code="MC-1", barcode="001")], [source(code="MC1", barcode="001")])
        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รหัสสินค้า"]), {"MC-1", "MC1"})

    def test_barcode_changes_do_not_change_calculations(self):
        a = [source(barcode="001")]
        g = [source(barcode="001", sales=12, stock=4)]
        before = combine(a, g)
        g[0]["barcode"] = "999"
        after = combine(a, g)
        columns = [c for c in before.columns if "barcode" not in c]
        pd.testing.assert_frame_equal(before[columns], after[columns])

    def test_barcode_priority_and_missing_sources(self):
        cases = [
            ([source(barcode="001")], [source(barcode="002")], "002"),
            ([source(barcode="001")], [source(barcode="")], "001"),
            ([source(barcode="001")], [], "001"),
            ([], [source(barcode="002")], "002"),
            ([source(barcode="")], [source(barcode="")], ""),
        ]
        for a, g, expected in cases:
            with self.subTest(expected=expected, asia=bool(a), green=bool(g)):
                self.assertEqual(combine(a, g).iloc[0]["barcode"], expected)

    def generate(self, tmp, catalog_rows):
        root = Path(tmp)
        catalog = root / "catalog.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "1G022"
        ws.append(["item", "image", "description", "brand", "material", "weight", "carton", "price", "barcode"])
        for row in catalog_rows:
            ws.append(row)
        wb.save(catalog)
        with patch.object(main, "PO_OUTPUT_FOLDER", tmp):
            output = main.generate_po_from_combined(
                combine([], [source(barcode="0000123")]), "1G022", datetime.date(2026, 10, 2),
                6, str(TEMPLATE), str(catalog), str(root / "vendors.xlsx"), 4, 7,
            )
        return openpyxl.load_workbook(output)["PO"]

    def test_po_keeps_march_columns_catalog_description_and_last_catalog_row(self):
        with tempfile.TemporaryDirectory() as tmp:
            ws = self.generate(tmp, [
                ["MC-1", None, "First catalog row", "Brand 1", "Steel", 1, 12, None, "0000123"],
                ["MC-1", None, "Last catalog row", "Brand 2", "Brass", 2, 24, None, "different"],
            ])
            self.assertEqual(ws["C9"].value, "Last catalog row")
            self.assertEqual(ws["D8"].value, "BRAND")
            self.assertEqual(ws["D9"].value, "Brand 2")
            self.assertEqual(ws["G9"].value, 24)
            self.assertEqual(ws["H9"].value, "=ROUND((Q9-U9)/G9,0)")
            self.assertEqual(ws["N9"].value, "=L9*K9")
            self.assertEqual(ws["X8"].value, "BARCODE")
            self.assertEqual(ws["X9"].value, "0000123")
            self.assertEqual(ws["X9"].data_type, "s")
            self.assertEqual(ws["X9"].number_format, "@")
            self.assertNotIn("หมายเหตุ", [c.value for c in ws[8]])
            self.assertIn("$X", str(ws.print_area))

    def test_missing_catalog_retains_march_blank_fields_and_unguarded_formulas(self):
        with tempfile.TemporaryDirectory() as tmp:
            ws = self.generate(tmp, [])
            self.assertIsNone(ws["G9"].value)
            self.assertEqual(ws["H9"].value, "=ROUND((Q9-U9)/G9,0)")
            self.assertEqual(ws["X9"].value, "0000123")


if __name__ == "__main__":
    unittest.main()
