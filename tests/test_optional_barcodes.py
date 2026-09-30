"""Behavioral regressions for optional barcodes and product variant fallbacks."""

from __future__ import annotations

import datetime
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pandas as pd

import main


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "ตัวอย่างใบสั่งซื้อต่างประเทศ.xlsx"


def source_row(
    description: str,
    *,
    code: str = "K1900",
    barcode: str = "",
    buyer: str = "A0029",
    sales: float = 0,
    stock: float = 0,
    on_order: float = 0,
) -> dict:
    return {
        "buyer": buyer,
        "รหัสสินค้า": code,
        "รายละเอียดสินค้า": description,
        "barcode": barcode,
        "ยอดขาย": sales,
        "สินค้าคงเหลือ": stock,
        "ON_ORDER": on_order,
        "หยวน": 10,
    }


def combine(asia: list[dict], green: list[dict]) -> pd.DataFrame:
    return main.build_combined_all(
        pd.DataFrame(asia), pd.DataFrame(green), months=3,
        min_factor=4, max_factor=7,
    )


def catalog_entry(description: str, brand: str, *, barcode: str = "", carton: int = 12) -> dict:
    return {
        "goods_desc": description,
        "brand": brand,
        "material": "Plastic",
        "weight": 1,
        "qty_per_carton": carton,
        "barcode": barcode,
        "img_bytes": None,
    }


class OptionalBarcodeTests(unittest.TestCase):
    def test_missing_barcodes_join_different_base_wording_and_sum_every_quantity(self) -> None:
        rows = combine([], [
            source_row("ฝารองนั่งวงรี", sales=2, stock=3, on_order=4),
            source_row("ฝารองนั่งรุ่นวงรี /OEM", sales=5, stock=6, on_order=7),
        ])

        self.assertEqual(len(rows), 1)
        row = rows.iloc[0]
        self.assertEqual(row["ยอดขาย_GREEN"], 7)
        self.assertEqual(row["STOCK_GREEN"], 9)
        self.assertEqual(row["ON_ORDER_GREEN"], 11)
        self.assertEqual(row["TOTAL_QTY_NUM"], 20)
        self.assertEqual(row["USE_MONTH"], 2)
        self.assertEqual(row["barcode"], "")

    def test_missing_barcode_identity_normalizes_buyer_and_item_code(self) -> None:
        rows = combine(
            [source_row("ก๊อกอ่างล้างหน้า /IR", code="A-1701-C", buyer=" a0029 ", sales=2)],
            [source_row("ก๊อกอ่างล้างหน้ารุ่นใหม่", code="A1701C", sales=3)],
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows.iloc[0]["buyer"], "A0029")
        self.assertEqual(rows.iloc[0]["รหัสสินค้า"], "A1701C")
        self.assertEqual(rows.iloc[0]["ยอดขาย_TOTAL"], 5)

    def test_source_markers_after_spaces_slashes_dashes_and_commas_do_not_duplicate_rows(self) -> None:
        for suffix in (" IR", "-MR", "/FN", ",VN", " /OEM", " /IR-MR,FN /VN-OEM"):
            with self.subTest(suffix=suffix):
                green_description = "ฝารองนั่งรุ่นวงรี-สีขาว" + suffix
                rows = combine(
                    [source_row("ฝารองนั่งวงรี-ขาว", sales=2)],
                    [source_row(green_description, sales=3)],
                )
                self.assertEqual(len(rows), 1)
                self.assertEqual(rows.iloc[0]["ยอดขาย_TOTAL"], 5)
                self.assertEqual(rows.iloc[0]["รายละเอียดสินค้า"], green_description)

    def test_white_variants_join_without_changing_full_green_description(self) -> None:
        green_description = "ฝารองนั่งรุ่นวงรี-สีขาว /OEM"
        rows = combine(
            [source_row("ฝารองนั่งวงรี-ขาว", sales=2, stock=1)],
            [source_row(green_description, sales=3, stock=2)],
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows.iloc[0]["รายละเอียดสินค้า"], green_description)
        self.assertEqual(rows.iloc[0]["ยอดขาย_TOTAL"], 5)
        self.assertEqual(rows.iloc[0]["TOTAL_QTY_NUM"], 3)

    def test_colors_after_a_space_remain_separate_without_barcodes(self) -> None:
        descriptions = (
            "ฝารองนั่งรุ่นวงรี สีขาว /OEM",
            "ฝารองนั่งรุ่นวงรี สีแดง /OEM",
            "ฝารองนั่งรุ่นวงรี สีฟ้า /OEM",
            "ฝารองนั่งรุ่นวงรี สีดำ /OEM",
        )
        rows = combine([], [source_row(description, sales=3) for description in descriptions])

        self.assertEqual(len(rows), 4)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), set(descriptions))
        self.assertEqual(rows["ยอดขาย_TOTAL"].sum(), 12)

    def test_thai_black_and_red_of_the_same_size_remain_separate(self) -> None:
        descriptions = ("สายอ่อนฝักบัว สีดำ 120CM", "สายอ่อนฝักบัว สีแดง 120CM")
        rows = combine([], [
            source_row(description, code="SHOWER", sales=3)
            for description in descriptions
        ])

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), set(descriptions))

    def test_thai_blue_and_brown_remain_separate_after_unicode_normalization(self) -> None:
        descriptions = ("พรมดักฝุ่น สีน้ำเงิน 30x40", "พรมดักฝุ่น สีน้ำตาล 30x40")
        rows = combine([], [
            source_row(description, code="FMR", sales=3)
            for description in descriptions
        ])

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), set(descriptions))

    def test_thai_cream_and_white_remain_separate(self) -> None:
        descriptions = ("ประตู PVC สีครีม", "ประตู PVC สีขาว")
        rows = combine([], [
            source_row(description, code="DP1-718CD", sales=3)
            for description in descriptions
        ])

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), set(descriptions))

    def test_light_dark_and_pastel_shades_remain_separate(self) -> None:
        descriptions = (
            "ก๊อก สีแดงอ่อน", "ก๊อก สีแดงเข้ม", "ก๊อก สีแดงพาสเทล", "ก๊อก สีแดง",
        )
        rows = combine([], [source_row(description, sales=3) for description in descriptions])

        self.assertEqual(len(rows), 4)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), set(descriptions))

    def test_shared_lb_marker_joins_blue_and_light_blue_wording(self) -> None:
        green_description = "ก็อกแฟนซีสี-ฟ้า(LB) /OEM"
        rows = combine(
            [source_row("ก็อกแฟนซีสี-น้ำเงิน(LB)", code="MC-510", sales=2)],
            [source_row(green_description, code="MC-510", sales=3)],
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows.iloc[0]["รายละเอียดสินค้า"], green_description)
        self.assertEqual(rows.iloc[0]["ยอดขาย_TOTAL"], 5)

    def test_short_and_long_handles_remain_separate_without_barcodes(self) -> None:
        descriptions = (
            "ก๊อกบอลด้ามสั้น 1/2 /FN",
            "ก๊อกบอลด้ามยาว 1/2 /FN",
        )
        rows = combine([], [
            source_row(description, code="AN105", sales=3)
            for description in descriptions
        ])

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), set(descriptions))

    def test_dimensions_remain_separate_without_barcodes(self) -> None:
        descriptions = ("อ่างล้างจาน 36x36 /IR", "อ่างล้างจาน 36x60 /IR")
        rows = combine([], [
            source_row(description, code="CE", sales=3)
            for description in descriptions
        ])

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), set(descriptions))

    def test_fraction_and_mixed_fraction_sizes_remain_separate_without_barcodes(self) -> None:
        descriptions = ("วาล์วทองเหลือง 1/2 /VN", "วาล์วทองเหลือง 1 1/4 /VN")
        rows = combine([], [
            source_row(description, code="NR197", sales=3)
            for description in descriptions
        ])

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), set(descriptions))

    def test_parenthetical_model_codes_remain_separate_without_barcodes(self) -> None:
        rows = combine([], [
            source_row("สายฉีดชำระรุ่นใหม่", code="NR56(LT)", sales=2),
            source_row("สายฉีดชำระรุ่นใหม่", code="NR56(AC)", sales=3),
        ])

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รหัสสินค้า"]), {"NR56(LT)", "NR56(AC)"})

    def test_missing_variant_detail_joins_one_compatible_cross_source_variant(self) -> None:
        rows = combine(
            [source_row("ฝารองนั่งวงรี", sales=2, stock=1)],
            [source_row("ฝารองนั่งรุ่นวงรี สีขาว", barcode="0000123", sales=3, stock=2)],
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows.iloc[0]["barcode"], "0000123")
        self.assertEqual(rows.iloc[0]["ยอดขาย_TOTAL"], 5)
        self.assertEqual(rows.iloc[0]["TOTAL_QTY_NUM"], 3)

    def test_missing_variant_detail_does_not_fan_out_to_two_colors(self) -> None:
        rows = combine(
            [source_row("ฝารองนั่งวงรี", sales=2, stock=1)],
            [
                source_row("ฝารองนั่งรุ่นวงรี สีขาว", barcode="0000123", sales=3, stock=2),
                source_row("ฝารองนั่งรุ่นวงรี สีแดง", barcode="0000456", sales=4, stock=3),
            ],
        )

        self.assertEqual(len(rows), 3)
        self.assertEqual(rows["ยอดขาย_TOTAL"].sum(), 9)
        self.assertEqual(rows["STOCK_ASIA"].sum(), 1)
        self.assertEqual(rows["STOCK_GREEN"].sum(), 5)
        self.assertEqual(rows.loc[rows["barcode"].ne(""), "ยอดขาย_ASIA"].sum(), 0)

    def test_generic_asia_row_is_not_completed_before_both_sources_are_considered(self) -> None:
        rows = combine(
            [
                source_row("ฝารองนั่งวงรี", sales=5),
                source_row("ฝารองนั่งวงรี สีแดง", sales=2),
            ],
            [
                source_row("ฝารองนั่งรุ่นวงรี สีแดง", barcode="0000123", sales=3),
                source_row("ฝารองนั่งรุ่นวงรี สีขาว", barcode="0000456", sales=4),
            ],
        )

        self.assertEqual(len(rows), 3)
        red = rows.loc[rows["barcode"].eq("0000123")].iloc[0]
        white = rows.loc[rows["barcode"].eq("0000456")].iloc[0]
        generic = rows.loc[rows["barcode"].eq("")].iloc[0]
        self.assertEqual((red["ยอดขาย_ASIA"], red["ยอดขาย_GREEN"]), (2, 3))
        self.assertEqual((white["ยอดขาย_ASIA"], white["ยอดขาย_GREEN"]), (0, 4))
        self.assertEqual((generic["ยอดขาย_ASIA"], generic["ยอดขาย_GREEN"]), (5, 0))
        self.assertEqual(rows["ยอดขาย_TOTAL"].sum(), 14)

    def test_redundant_decimal_places_do_not_create_another_dimension_variant(self) -> None:
        rows = combine(
            [source_row('พรมดักฝุ่น 15.50x23.50"', code="FDM", sales=2)],
            [source_row('พรมดักฝุ่น 15.5x23.5"', code="FDM", sales=3)],
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows.iloc[0]["ยอดขาย_TOTAL"], 5)

    def test_uncoded_products_with_different_names_remain_separate(self) -> None:
        rows = combine(
            [source_row("ก๊อกอ่างล้างหน้า", code="", sales=2)],
            [source_row("ฝารองนั่งวงรี", code="", sales=3)],
        )

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), {"ก๊อกอ่างล้างหน้า", "ฝารองนั่งวงรี"})
        self.assertEqual(rows["ยอดขาย_TOTAL"].sum(), 5)

    def test_welcome_small_and_large_without_numeric_sizes_remain_separate(self) -> None:
        descriptions = ("ยางปูหน้าอาคาร ขนาดเล็ก", "ยางปูหน้าอาคาร ขนาดใหญ่")
        rows = combine([], [
            source_row(description, code="WELCOME", sales=3)
            for description in descriptions
        ])

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), set(descriptions))

    def test_round_and_square_drain_cups_remain_separate(self) -> None:
        descriptions = ("ถ้วยกันกลิ่นสแตนเลส กลม", "ถ้วยกันกลิ่นสแตนเลส เหลี่ยม")
        rows = combine([], [
            source_row(description, code="DM-2114", sales=3)
            for description in descriptions
        ])

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), set(descriptions))

    def test_rain_letter_models_remain_separate_with_identical_head_shape(self) -> None:
        descriptions = ("รุ่น R-หัว RAIN กลม", "รุ่น L-หัว RAIN กลม")
        rows = combine([], [
            source_row(description, code="RAIN", sales=3)
            for description in descriptions
        ])

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), set(descriptions))

    def test_same_cream_color_does_not_merge_head_only_and_complete_set(self) -> None:
        descriptions = ("ฝักบัวชำระ สีครีม หัวอย่างเดียว", "ฝักบัวชำระ สีครีม ครบชุด")
        rows = combine([], [
            source_row(description, code="CREAM", sales=3)
            for description in descriptions
        ])

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), set(descriptions))

    def test_same_sink_with_legs_and_without_legs_remain_separate(self) -> None:
        descriptions = ("อ่างล้างหน้า พร้อมขา", "อ่างล้างหน้า ไม่มีขา")
        rows = combine([], [
            source_row(description, code="CBS-320", sales=3)
            for description in descriptions
        ])

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), set(descriptions))

    def test_truncated_brown_and_blue_color_words_remain_separate(self) -> None:
        descriptions = ("พรมดักฝุ่น สีน้ำต", "พรมดักฝุ่น สีน้ำเง")
        rows = combine([], [
            source_row(description, code="FDM", sales=3)
            for description in descriptions
        ])

        self.assertEqual(len(rows), 2)
        self.assertEqual(set(rows["รายละเอียดสินค้า"]), set(descriptions))

    def test_different_explicit_colors_cannot_join_cross_source(self) -> None:
        rows = combine(
            [source_row("ฝารองนั่งวงรี สีขาว", sales=2)],
            [source_row("ฝารองนั่งรุ่นวงรี สีแดง", barcode="0000456", sales=3)],
        )

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows["ยอดขาย_TOTAL"].sum(), 5)

    def test_asia_barcode_does_not_populate_a_missing_green_output_barcode(self) -> None:
        rows = combine(
            [source_row("ฝารองนั่งวงรี-ขาว", barcode="0000123", sales=2)],
            [source_row("ฝารองนั่งรุ่นวงรี-สีขาว", sales=3)],
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows.iloc[0]["barcode"], "")
        self.assertEqual(rows.iloc[0]["catalog_match_barcode"], "0000123")

    def test_sole_catalog_code_is_allowed_despite_multiple_source_variants(self) -> None:
        catalog = {"K1900": [catalog_entry("Old catalog wording", "Sole row")]}

        result = main.resolve_catalog_variant(
            catalog, "K1900", "ฝารองนั่งวงรี-ขาว", 7,
            description_variant_count=3, barcode="",
        )

        self.assertEqual(result["brand"], "Sole row")
        self.assertEqual(result["qty_per_carton"], 12)

    def test_missing_source_barcode_uses_last_exact_description_match(self) -> None:
        catalog = {"K1900": [
            catalog_entry("ฝารองนั่งวงรี-ขาว", "First exact", carton=12),
            catalog_entry("ฝารองนั่งวงรี-ขาว", "Last exact", carton=24),
            catalog_entry("ฝารองนั่งวงรี-แดง", "Later other color", carton=48),
        ]}

        result = main.resolve_catalog_variant(
            catalog, "K1900", "ฝารองนั่งวงรี-ขาว", 3,
            description_variant_count=2, barcode="",
        )

        self.assertEqual(result["brand"], "Last exact")
        self.assertEqual(result["qty_per_carton"], 24)

    def test_missing_catalog_barcode_prefers_last_matching_color_over_last_code_row(self) -> None:
        catalog = {"K1900": [
            catalog_entry("ฝารองนั่งวงรี-ขาว", "First white", carton=12),
            catalog_entry("ฝารองนั่งรุ่นวงรี สีขาว", "Last white", carton=24),
            catalog_entry("ฝารองนั่งวงรี สีแดง", "Last code row", carton=48),
        ]}

        result = main.resolve_catalog_variant(
            catalog, "K1900", "ฝารองนั่งวงรี-สีขาว /OEM", 3,
            barcode="0000123",
        )

        self.assertEqual(result["brand"], "Last white")
        self.assertEqual(result["qty_per_carton"], 24)

    def test_missing_source_barcode_uses_march_last_code_row_if_no_description_match(self) -> None:
        catalog = {"K1900": [
            catalog_entry("Old first description", "First", barcode="0000123", carton=12),
            catalog_entry("Old last description", "Last", barcode="0000456", carton=24),
        ]}

        result = main.resolve_catalog_variant(
            catalog, "K1900", "New source wording", 4, barcode="",
        )

        self.assertEqual(result["brand"], "Last")
        self.assertEqual(result["qty_per_carton"], 24)

    def test_unknown_source_barcode_in_partial_catalog_falls_back_only_to_blank_rows(self) -> None:
        catalog = {"K1900": [
            catalog_entry("Blank old first description", "First eligible", carton=12),
            catalog_entry("Blank old last description", "Last eligible", carton=24),
            catalog_entry("New source wording", "Conflicting barcode", barcode="0000999", carton=48),
        ]}

        result = main.resolve_catalog_variant(
            catalog, "K1900", "New source wording", 4, barcode="0000123",
        )

        self.assertEqual(result["brand"], "Last eligible")
        self.assertEqual(result["qty_per_carton"], 24)

    def test_populated_conflicting_catalog_barcodes_are_not_optional_barcode_fallbacks(self) -> None:
        catalog = {"K1900": [
            catalog_entry("New source wording", "Conflicting barcode", barcode="0000999"),
        ]}

        with self.assertRaises(ValueError):
            main.resolve_catalog_variant(
                catalog, "K1900", "New source wording", 1, barcode="0000123",
            )

    def test_unique_catalog_barcode_has_priority_over_description_and_code_fallbacks(self) -> None:
        catalog = {
            "K1900": [catalog_entry("Source description", "Code-only fallback", carton=12)],
            "TYPO-CODE": [catalog_entry("Different catalog wording", "Barcode match", barcode="0000123", carton=24)],
        }

        result = main.resolve_catalog_variant(
            catalog, "K1900", "Source description", 3, barcode="0000123",
        )

        self.assertEqual(result["brand"], "Barcode match")
        self.assertEqual(result["qty_per_carton"], 24)

    def test_generated_po_keeps_full_green_description_and_final_text_barcode(self) -> None:
        green_description = "ฝารองนั่งรุ่นวงรี สีขาว /OEM"
        rows = combine(
            [source_row("ฝารองนั่งวงรี-ขาว", sales=6, stock=1)],
            [source_row(green_description, barcode="0000123", sales=9, stock=2)],
        )
        self.assertEqual(len(rows), 1)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            catalog_path = output / "optional_barcodes.xlsx"
            wb = openpyxl.Workbook()
            catalog_ws = wb.active
            catalog_ws.title = "A0029"
            catalog_ws.append([
                "BUYER ITEM NO.", "GOODS PICTURE", "GOODS DESCRIPTION", "BRAND",
                "MATERIAL", "Weight", "QTY PER CARTON", "UNIT PRICE (FOB)", "BARCODE",
            ])
            catalog_ws.append(["K1900", None, "ฝารองนั่งวงรี-ขาว", "Donmark", "Plastic", 1, 12, None, None])
            wb.save(catalog_path)

            with patch.object(main, "PO_OUTPUT_FOLDER", tmp):
                path = main.generate_po_from_combined(
                    rows, "A0029", datetime.date(2026, 9, 30), 6,
                    str(TEMPLATE), str(catalog_path), str(output / "missing_vendors.xlsx"),
                    4, 7,
                )

            po = openpyxl.load_workbook(path)["PO"]
            columns = main.get_po_col_map(po)
            self.assertEqual(po.cell(9, columns["BUYER ITEM NO."]).value, "K1900")
            self.assertEqual(po.cell(9, columns["GOODS DESCRIPTION"]).value, green_description)
            self.assertEqual(po.cell(9, columns["QTY PER CARTON"]).value, 12)
            self.assertEqual(po.cell(9, columns["STOCK ASIA"]).value, 1)
            self.assertEqual(po.cell(9, columns["STOCK GREEN"]).value, 2)
            self.assertEqual(po.cell(9, columns["USE MONTH"]).value, 5)
            barcode_cell = po.cell(9, columns["BARCODE"])
            self.assertEqual(barcode_cell.value, "0000123")
            self.assertEqual(barcode_cell.data_type, "s")
            self.assertEqual(barcode_cell.number_format, "@")
            self.assertEqual(columns["BARCODE"], max(columns.values()))


if __name__ == "__main__":
    unittest.main()
