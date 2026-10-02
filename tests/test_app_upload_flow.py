"""Exercise Streamlit reruns and upload confirmation through the actual app.

AppTest does not expose a file-uploader setter. Only that browser boundary is
stubbed; workbook detection, rendered widgets, button events and session state
all run through Streamlit. PO generation is captured at its public boundary.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import unittest
import inspect
from unittest.mock import patch

import openpyxl
from streamlit.testing.v1 import AppTest

import main


ROOT = Path(__file__).resolve().parents[1]
HEADINGS = {
    "ASIA": "บริษัท เอเซีย โฮม สแตนดาร์ด จำกัด",
    "GREEN": "บริษัท กรีนไลฟ์ เอ็นเตอร์ไพรส์ จำกัด",
}


class UploadedWorkbook(BytesIO):
    def __init__(self, name: str, data: bytes):
        super().__init__(data)
        self.name = name
        self.size = len(data)


def workbook_upload(name: str, heading: str, *, marker: str = "") -> UploadedWorkbook:
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = heading
    workbook.active["A10"] = marker
    data = BytesIO()
    workbook.save(data)
    return UploadedWorkbook(name, data.getvalue())


def company_upload(source: str, *, name: str | None = None, marker: str = "") -> UploadedWorkbook:
    return workbook_upload(name or f"{source}.xlsx", HEADINGS[source], marker=marker)


class AppUploadFlowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.uploads: list[UploadedWorkbook] = []
        self.catalog = workbook_upload("รายละเอียดสินค้า ตุลาคม 2569.xlsx", "catalog")
        self.vendor_info = workbook_upload("ข้อมูลผู้จำหน่าย.xlsx", "vendors")
        self.generated: list[dict] = []
        uploader_patch = patch("streamlit.file_uploader", side_effect=self.file_uploader)
        self.generation_signature = inspect.signature(main.generate_po_streamlit)
        generation_patch = patch.object(main, "generate_po_streamlit", side_effect=self.generate)
        uploader_patch.start()
        generation_patch.start()
        self.addCleanup(uploader_patch.stop)
        self.addCleanup(generation_patch.stop)
        self.app = AppTest.from_file(ROOT / "app.py", default_timeout=10)

    def file_uploader(self, label: str, *args, key: str | None = None, **kwargs):
        if key == "express":
            self.assertTrue(kwargs.get("accept_multiple_files"))
            return self.uploads
        if key == "catalog":
            return self.catalog
        if key == "vendorinfo":
            return self.vendor_info
        self.fail(f"Unexpected upload field: {key!r}, {label!r}")

    def generate(self, **kwargs) -> dict:
        self.generation_signature.bind(**kwargs)
        capture = dict(kwargs)
        for source, argument in (("ASIA", "express_asia_path"), ("GREEN", "express_green_path")):
            path = kwargs[argument]
            capture[f"{source}_bytes"] = Path(path).read_bytes() if path else None
        capture["catalog_bytes"] = Path(kwargs["catalog_path"]).read_bytes()
        capture["vendor_bytes"] = Path(kwargs["vendor_info_path"]).read_bytes()
        self.generated.append(capture)
        output = Path(kwargs["catalog_path"]).parent / "generated_all_items.xlsx"
        output.write_bytes(self.catalog.getvalue())
        return {
            "po_all_items": str(output),
            "po_filtered": None,
            "count_all": 1,
            "count_filtered": 0,
        }

    def run_app(self) -> None:
        self.app.run()
        self.assertFalse(self.app.exception, [error.message for error in self.app.exception])

    def set_supplier(self) -> None:
        self.app.text_input[0].set_value("a0029").run()
        self.assertFalse(self.app.exception)

    def confirm_single(self) -> None:
        self.app.button(key="confirm_single_express").click().run()
        self.assertFalse(self.app.exception)
        self.assertFalse(self.app.button(key="generate_po").disabled)

    def generate_po(self) -> None:
        self.app.button(key="generate_po").click().run()
        self.assertFalse(self.app.exception)
        self.assertFalse(self.app.error, [error.value for error in self.app.error])

    def test_no_express_files_disable_generation(self) -> None:
        self.run_app()
        self.assertTrue(self.app.button(key="generate_po").disabled)
        self.assertNotIn("confirm_single_express", [button.key for button in self.app.button])
        self.assertEqual(self.generated, [])

    def test_each_single_company_requires_confirmation_and_routes_only_present_source(self) -> None:
        for source in ("ASIA", "GREEN"):
            with self.subTest(source=source):
                self.generated.clear()
                self.app = AppTest.from_file(ROOT / "app.py", default_timeout=10)
                upload = company_upload(source, name="generic Express export.xlsx")
                self.uploads = [upload]
                self.run_app()
                self.set_supplier()
                self.assertTrue(self.app.button(key="generate_po").disabled)
                self.assertEqual(self.generated, [])
                self.confirm_single()
                self.generate_po()
                self.assertEqual(len(self.generated), 1)
                capture = self.generated[-1]
                absent = "GREEN" if source == "ASIA" else "ASIA"
                self.assertEqual(capture[f"{source}_bytes"], upload.getvalue())
                self.assertIsNone(capture[f"{absent}_bytes"])
                absent_argument = "express_green_path" if source == "ASIA" else "express_asia_path"
                self.assertEqual(capture[absent_argument], "")
                self.assertEqual(capture["vendor_code"], "A0029")
                self.assertEqual(capture["catalog_bytes"], self.catalog.getvalue())
                self.assertEqual(capture["vendor_bytes"], self.vendor_info.getvalue())

    def test_two_company_files_work_in_either_order_despite_misleading_filenames(self) -> None:
        asia = company_upload("ASIA", name="GREEN.xlsx")
        green = company_upload("GREEN", name="ASIA.xlsx")
        for uploads in ([asia, green], [green, asia]):
            with self.subTest(order=[upload.name for upload in uploads]):
                self.generated.clear()
                self.app = AppTest.from_file(ROOT / "app.py", default_timeout=10)
                self.uploads = uploads
                self.run_app()
                self.set_supplier()
                self.assertFalse(self.app.button(key="generate_po").disabled)
                self.assertNotIn("confirm_single_express", [button.key for button in self.app.button])
                captions = "\n".join(caption.value for caption in self.app.caption)
                self.assertIn(asia.name, captions)
                self.assertIn(green.name, captions)
                self.generate_po()
                self.assertEqual(len(self.generated), 1)
                capture = self.generated[-1]
                self.assertEqual(capture["ASIA_bytes"], asia.getvalue())
                self.assertEqual(capture["GREEN_bytes"], green.getvalue())

    def test_confirmation_survives_unrelated_input_reruns(self) -> None:
        self.uploads = [company_upload("GREEN")]
        self.run_app()
        self.confirm_single()
        self.set_supplier()
        self.app.number_input[0].set_value(5).run()
        self.assertFalse(self.app.exception)
        self.assertFalse(self.app.button(key="generate_po").disabled)
        self.generate_po()
        self.assertEqual(self.generated[-1]["min_factor"], 5)

    def test_same_filename_with_changed_content_requires_new_confirmation(self) -> None:
        self.uploads = [company_upload("GREEN", name="Express.xlsx", marker="original")]
        self.run_app()
        self.confirm_single()
        self.uploads = [company_upload("GREEN", name="Express.xlsx", marker="replaced")]
        self.run_app()
        self.assertTrue(self.app.button(key="generate_po").disabled)
        self.assertEqual(self.generated, [])
        self.confirm_single()

    def test_renamed_single_upload_requires_new_confirmation(self) -> None:
        upload = company_upload("ASIA", name="old.xlsx")
        self.uploads = [upload]
        self.run_app()
        self.confirm_single()
        self.uploads = [UploadedWorkbook("renamed.xlsx", upload.getvalue())]
        self.run_app()
        self.assertTrue(self.app.button(key="generate_po").disabled)
        self.assertEqual(self.generated, [])

    def test_removing_and_readding_the_same_file_requires_new_confirmation(self) -> None:
        upload = company_upload("GREEN")
        self.uploads = [upload]
        self.run_app()
        self.confirm_single()
        self.uploads = []
        self.run_app()
        self.assertTrue(self.app.button(key="generate_po").disabled)
        self.uploads = [upload]
        self.run_app()
        self.assertTrue(self.app.button(key="generate_po").disabled)
        self.assertEqual(self.generated, [])

    def test_single_to_pair_to_single_does_not_reuse_confirmation(self) -> None:
        green = company_upload("GREEN")
        self.uploads = [green]
        self.run_app()
        self.confirm_single()
        self.uploads = [green, company_upload("ASIA")]
        self.run_app()
        self.assertFalse(self.app.button(key="generate_po").disabled)
        self.uploads = [green]
        self.run_app()
        self.assertTrue(self.app.button(key="generate_po").disabled)
        self.assertEqual(self.generated, [])

    def test_duplicate_unknown_ambiguous_and_excess_uploads_block_generation(self) -> None:
        invalid_sets = {
            "duplicate company": [company_upload("GREEN", name="one.xlsx"), company_upload("GREEN", name="two.xlsx")],
            "unknown company": [workbook_upload("GREEN.xlsx", "Unrecognized Company Limited")],
            "ambiguous company": [workbook_upload("both.xlsx", HEADINGS["ASIA"] + " " + HEADINGS["GREEN"])],
            "unreadable workbook": [UploadedWorkbook("broken.xlsx", b"not an Excel workbook")],
            "too many files": [company_upload("ASIA"), company_upload("GREEN"), company_upload("GREEN", name="extra.xlsx")],
        }
        for label, uploads in invalid_sets.items():
            with self.subTest(case=label):
                self.app = AppTest.from_file(ROOT / "app.py", default_timeout=10)
                self.uploads = uploads
                self.run_app()
                self.assertTrue(self.app.button(key="generate_po").disabled)
                self.assertTrue(self.app.error)
                self.assertNotIn("confirm_single_express", [button.key for button in self.app.button])
                self.assertEqual(self.generated, [])


if __name__ == "__main__":
    unittest.main()
